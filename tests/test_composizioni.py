import io

import pytest
from docx import Document

from sicurcantiere import composizioni as comp
from sicurcantiere.archivio import LAVORAZIONI
from sicurcantiere.db import Progetto
from sicurcantiere.documenti import composizione
from sicurcantiere.repository import Repository


@pytest.fixture
def repo(tmp_path):
    progetto = Progetto(tmp_path / "prova.spo", crea=True)
    r = Repository(progetto)
    r.aggiorna_cantiere(nome="Villa Rossi", comune="Bari", indirizzo="Via Roma 1")
    impresa = r.inserisci("imprese", ragione_sociale="Edil Alfa srl", datore_lavoro="Mario Bianchi")
    r.importa_lavorazione(next(l for l in LAVORAZIONI if l["codice"] == "AL.21"))
    for f in r.elenco("fasi"):
        r.aggiorna("fasi", f["id"], impresa_id=impresa)
    r.impresa = impresa
    yield r
    progetto.chiudi()


def _testo(dati: bytes) -> str:
    d = Document(io.BytesIO(dati))
    parti = [p.text for p in d.paragraphs]
    parti += [c.text for t in d.tables for riga in t.rows for c in riga.cells]
    return "\n".join(parti)


def _xml_intestazioni(dati: bytes) -> str:
    d = Document(io.BytesIO(dati))
    return "".join(s.header._element.xml + s.first_page_header._element.xml for s in d.sections
                   if not s.header.is_linked_to_previous)


def test_bozza_con_dati_mancanti_e_revisione(repo):
    documento = f"POS:{repo.impresa}"
    c = comp.componi(repo, documento)
    assert c["stato"] == comp.BOZZA and c["numero"] == 0 and c["descrizione"] == "Prima emissione"
    revisioni = repo.elenco("revisioni", documento=documento)
    assert [(r["numero"], r["descrizione"]) for r in revisioni] == [(0, "Prima emissione")]

    dati_mancanti = {m["dato"]: m for m in c["mancanti"]}
    assert dati_mancanti["committente"]["livello"] == composizione.OBBLIGATORIO
    assert dati_mancanti["nominativo: RSPP"]["sezione"] == "Figure della sicurezza"
    assert dati_mancanti["nominativo: RSPP"]["capitolo"] == "4 Organizzazione della sicurezza"
    assert dati_mancanti["titolo abilitativo"]["livello"] == composizione.FACOLTATIVO

    dati = comp.contenuto(repo, c["id"])
    testo = _testo(dati)
    assert "[Da compilare: committente]" in testo
    assert "[Da compilare: titolo abilitativo]" not in testo  # i facoltativi non si stampano
    assert "Dati da completare" in testo  # riepilogo della bozza
    assert "BOZZA" in _xml_intestazioni(dati)  # filigrana
    assert composizione.leggi_mancanti(dati) > 5
    assert any(m.author == composizione.AUTORE for m in Document(io.BytesIO(dati)).comments)

    # Ricomporre aggiorna la stessa bozza, non ne crea un'altra.
    repo.aggiorna("imprese", repo.impresa, telefono="080 123456")
    di_nuovo = comp.componi(repo, documento)
    assert di_nuovo["id"] == c["id"] and len(comp.elenco(repo, documento)) == 1
    assert "telefono dell'impresa" not in {m["dato"] for m in di_nuovo["mancanti"]}
    assert len(repo.elenco("revisioni", documento=documento)) == 1


def test_emissione_consegna_e_nuova_revisione(repo):
    documento = f"POS:{repo.impresa}"
    c = comp.componi(repo, documento)
    with pytest.raises(comp.ErroreComposizione):
        comp.emetti(repo, c["id"])  # dati obbligatori mancanti: serve la conferma
    emessa = comp.emetti(repo, c["id"], "Emessa con dati mancanti. Motivo: prova")
    assert emessa["stato"] == comp.EMESSA and emessa["emessa_il"] and "prova" in emessa["conferma"]
    dati = comp.contenuto(repo, c["id"])
    testo = _testo(dati)
    assert "Da compilare" not in testo and "Dati da completare" not in testo
    assert "BOZZA" not in _xml_intestazioni(dati)
    assert composizione.RIGA_PENNA in testo  # i dati mancanti diventano righe da compilare a mano
    assert composizione.leggi_mancanti(dati) == 0
    assert not any(m.author == composizione.AUTORE for m in Document(io.BytesIO(dati)).comments)
    with pytest.raises(comp.ErroreComposizione):
        comp.emetti(repo, c["id"], "x")  # una revisione emessa non si emette di nuovo

    comp.segna_consegnata(repo, c["id"], "CSE ing. Verdi", "2026-10-22")
    assert comp.leggi(repo, c["id"])["stato"] == comp.CONSEGNATA

    # La revisione successiva: nuova bozza con il numero seguente, nella tabella delle revisioni.
    assert comp.numero_proposto(repo, documento) == 1
    nuova = comp.componi(repo, documento, "Ingresso nuova lavorazione")
    assert nuova["id"] != c["id"] and nuova["numero"] == 1
    assert "Ingresso nuova lavorazione" in _testo(comp.contenuto(repo, nuova["id"]))
    assert [r["numero"] for r in repo.elenco("revisioni", documento=documento)] == [0, 1]

    # Eliminare la bozza toglie anche la sua revisione; annullare un'emessa la lascia nello storico.
    comp.annulla(repo, nuova["id"])
    assert comp.leggi(repo, nuova["id"]) is None
    assert [r["numero"] for r in repo.elenco("revisioni", documento=documento)] == [0]
    comp.annulla(repo, c["id"])
    assert comp.leggi(repo, c["id"])["stato"] == comp.ANNULLATA
    assert repo.elenco("revisioni", documento=documento)[0]["descrizione"].startswith("Annullata - ")


def test_capitoli_personalizzati_restano_e_confronto(repo):
    documento = f"POS:{repo.impresa}"
    prima = comp.componi(repo, documento)
    chiavi = {c["chiave"]: c for c in prima["capitoli"]}
    assert "documentazione-da-custodire-in-cantiere" in chiavi
    integrative = "misure-preventive-e-protettive-integrative-rispetto-al-psc"
    comp.salva_personalizzazione(repo, documento, "documentazione-da-custodire-in-cantiere", escluso=True)
    comp.salva_personalizzazione(repo, documento, integrative, modo=composizione.SOSTITUISCI,
                                 testo="Nessuna misura integrativa oltre al PSC.\n- verifica giornaliera",
                                 impronta=chiavi[integrative]["impronta"])
    comp.salva_personalizzazione(repo, documento, "segnaletica-prevista-nel-cantiere",
                                 testo="Cartelli forniti dall'impresa affidataria.")
    dopo = comp.componi(repo, documento)
    testo = _testo(comp.contenuto(repo, dopo["id"]))
    assert "Documentazione da custodire in cantiere" not in testo
    assert "Nessuna misura integrativa oltre al PSC." in testo and "verifica giornaliera" in testo
    assert "Cartelli forniti dall'impresa affidataria." in testo
    # sostituito il contenuto, il suo dato mancante non c'è più (il riferimento di legge resta)
    assert not any(m["dato"].startswith("misure integrative") for m in dopo["mancanti"])
    assert "Allegato XV punto 3.2.1, lettera g)" in testo
    voci = {c["chiave"]: c for c in dopo["capitoli"]}
    assert voci["documentazione-da-custodire-in-cantiere"]["escluso"]
    assert voci[integrative]["modo"] == composizione.SOSTITUISCI and not voci[integrative]["cambiato"]

    # Personalizzazione vuota: si toglie.
    comp.salva_personalizzazione(repo, documento, "segnaletica-prevista-nel-cantiere", testo="")
    assert "segnaletica-prevista-nel-cantiere" not in comp.personalizzazioni(repo, documento)

    cambiamenti = {c["titolo"]: c for c in comp.confronta(prima, dopo)}
    assert cambiamenti["Documentazione da custodire in cantiere"]["stato"] == "tolto"
    assert cambiamenti["Misure preventive e protettive integrative rispetto al PSC"]["stato"] == "cambiato"
    assert ("+", "Nessuna misura integrativa oltre al PSC.") in \
        cambiamenti["Misure preventive e protettive integrative rispetto al PSC"]["righe"]


def test_bozza_modificata_in_word_e_psc(repo, tmp_path):
    documento = f"POS:{repo.impresa}"
    c = comp.componi(repo, documento)
    d = Document(io.BytesIO(comp.contenuto(repo, c["id"])))
    d.add_paragraph("Aggiunto a mano in Word")
    uscita = io.BytesIO()
    d.save(uscita)
    comp.acquisisci_da_word(repo, c["id"], uscita.getvalue())
    assert comp.leggi(repo, c["id"])["modificata_word"] == 1
    emessa = comp.emetti(repo, c["id"], "conferma")
    testo = _testo(comp.contenuto(repo, emessa["id"]))
    assert "Aggiunto a mano in Word" in testo and "Da compilare" not in testo

    psc = comp.componi(repo, "PSC")
    assert psc["stato"] == comp.BOZZA and psc["capitoli"]
    assert "BOZZA" in _xml_intestazioni(comp.contenuto(repo, psc["id"]))
