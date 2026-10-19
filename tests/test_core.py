from datetime import date

import pytest
from docx import Document

from sicurcantiere import analisi
from sicurcantiere.archivio import LAVORAZIONI
from sicurcantiere.db import ErroreProgetto, Progetto, VERSIONE_SCHEMA
from sicurcantiere.documenti.psc import genera_psc
from sicurcantiere.repository import Repository


@pytest.fixture
def repo(tmp_path):
    progetto = Progetto(tmp_path / "prova.scs", crea=True)
    yield Repository(progetto)
    progetto.chiudi()


def test_nuovo_progetto_ha_schema_aggiornato(tmp_path):
    p = Progetto(tmp_path / "a.scs", crea=True)
    assert p.query("PRAGMA user_version")[0][0] == VERSIONE_SCHEMA
    p.chiudi()
    # Riapertura: crea il backup e non riapplica le migrazioni.
    p = Progetto(tmp_path / "a.scs")
    assert (tmp_path / "a.scs.bak").exists()
    p.chiudi()


def test_file_non_valido(tmp_path):
    f = tmp_path / "rotto.scs"
    f.write_bytes(b"non sono un database" * 100)
    with pytest.raises(ErroreProgetto):
        Progetto(f)


def test_versione_futura_rifiutata(tmp_path):
    p = Progetto(tmp_path / "f.scs", crea=True)
    p.conn.execute(f"PRAGMA user_version = {VERSIONE_SCHEMA + 1}")
    p.chiudi()
    with pytest.raises(ErroreProgetto):
        Progetto(tmp_path / "f.scs")


def test_campi_non_validi_rifiutati(repo):
    with pytest.raises(ValueError):
        repo.inserisci("imprese", **{"ragione_sociale; DROP TABLE imprese": "x"})
    with pytest.raises(ValueError):
        repo.elenco("sqlite_master")


def test_crud_e_cascata(repo):
    impresa = repo.inserisci("imprese", ragione_sociale="Edil Srl")
    fase = repo.importa_lavorazione(LAVORAZIONI[0])
    repo.aggiorna("fasi", fase, impresa_id=impresa)
    assert repo.leggi("fasi", fase)["impresa_id"] == impresa
    assert len(repo.elenco("rischi", fase_id=fase)) == len(LAVORAZIONI[0]["rischi"])
    repo.elimina("imprese", impresa)
    assert repo.leggi("fasi", fase)["impresa_id"] is None
    repo.elimina("fasi", fase)
    assert repo.elenco("rischi") == []


def test_salva_con_nome(repo, tmp_path):
    repo.aggiorna_cantiere(nome="Scuola")
    repo.progetto.salva_con_nome(tmp_path / "copia.scs")
    copia = Progetto(tmp_path / "copia.scs")
    assert Repository(copia).cantiere()["nome"] == "Scuola"
    copia.chiudi()


def test_classi_rischio():
    assert analisi.classe_rischio(analisi.livello_rischio(1, 2)) == "Basso"
    assert analisi.classe_rischio(analisi.livello_rischio(2, 2)) == "Medio"
    assert analisi.classe_rischio(analisi.livello_rischio(2, 3)) == "Alto"
    assert analisi.classe_rischio(analisi.livello_rischio(3, 4)) == "Molto alto"


def test_giorni_lavorativi():
    # 6-19 ottobre 2025: due settimane complete, lunedì -> domenica.
    assert analisi.giorni_lavorativi(date(2025, 10, 6), date(2025, 10, 19)) == 10
    assert analisi.giorni_lavorativi(date(2025, 10, 10), date(2025, 10, 13)) == 2  # ven -> lun
    assert analisi.giorni_lavorativi(date(2025, 10, 11), date(2025, 10, 12)) == 0
    assert analisi.giorni_lavorativi(date(2025, 10, 13), date(2025, 10, 10)) == 0


def _fase(id_, impresa, inizio, fine, lavoratori=1):
    return {"id": id_, "codice": f"F{id_}", "descrizione": "", "impresa_id": impresa,
            "data_inizio": inizio, "data_fine": fine, "lavoratori": lavoratori}


def test_interferenze():
    fasi = [
        _fase(1, 1, "2025-10-06", "2025-10-17"),
        _fase(2, 2, "2025-10-15", "2025-10-24"),  # si sovrappone a 1, altra impresa
        _fase(3, 1, "2025-10-16", "2025-10-16"),  # stessa impresa di 1, sovrapposta a 2
        _fase(4, 2, "2025-11-01", "2025-11-05"),  # isolata
        _fase(5, None, None, None),
    ]
    coppie = {(i.fase_a["id"], i.fase_b["id"]) for i in analisi.interferenze(fasi)}
    assert coppie == {(1, 2), (2, 3)}
    prima = analisi.interferenze(fasi)[0]
    assert (prima.inizio, prima.fine) == (date(2025, 10, 15), date(2025, 10, 17))


def test_obblighi():
    imprese = [{"id": 1, "tipo": "Affidataria"}, {"id": 2, "tipo": "Esecutrice"},
               {"id": 3, "tipo": "Lavoratore autonomo"}]
    una = [_fase(1, 1, "2025-10-06", "2025-10-10", 3)]
    o = analisi.obblighi(imprese, una)
    assert (o.imprese, o.uomini_giorno, o.psc, o.notifica) == (1, 15, False, False)

    grande = [_fase(1, 1, "2025-01-06", "2025-03-28", 4)]  # 60 gg lavorativi x 4
    o = analisi.obblighi(imprese, grande)
    assert o.uomini_giorno == 240 and not o.psc and o.notifica

    due = una + [_fase(2, 2, "2025-12-01", "2025-12-02")]
    o = analisi.obblighi(imprese, due)
    assert o.psc and o.notifica


def test_genera_psc(repo, tmp_path):
    repo.aggiorna_cantiere(nome="Ristrutturazione scuola", comune="Bari", provincia="BA",
                           importo_lavori=250000)
    repo.inserisci("soggetti", ruolo="Committente", nome="Comune di Bari")
    a = repo.inserisci("imprese", ragione_sociale="Alfa Costruzioni", tipo="Affidataria")
    b = repo.inserisci("imprese", ragione_sociale="Beta Impianti")
    f1 = repo.importa_lavorazione(LAVORAZIONI[2])
    f2 = repo.importa_lavorazione(LAVORAZIONI[6])
    repo.aggiorna("fasi", f1, impresa_id=a, data_inizio="2025-10-06", data_fine="2025-10-31")
    repo.aggiorna("fasi", f2, impresa_id=b, data_inizio="2025-10-20", data_fine="2025-11-14")
    repo.inserisci("costi", descrizione="Recinzione", unita="m", quantita=120, prezzo_unitario=12.5)
    repo.salva_testo("emergenze", "Presidio di primo soccorso nel box.")

    file = genera_psc(repo, tmp_path / "psc.docx")
    doc = Document(file)
    testo = "\n".join(p.text for p in doc.paragraphs)
    celle = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "PIANO DI SICUREZZA E COORDINAMENTO" in testo
    assert "Presidio di primo soccorso" in testo
    assert "Alfa Costruzioni" in celle
    assert "€ 1.500,00" in celle
    assert "Scavo a sezione aperta" in celle  # interferenza e cronoprogramma


def test_genera_psc_progetto_vuoto(repo, tmp_path):
    assert genera_psc(repo, tmp_path / "vuoto.docx").exists()
