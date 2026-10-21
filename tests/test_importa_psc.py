from docx import Document

from sicurcantiere import importa_psc
from sicurcantiere.ai import OPENAI, AssistenteAI
from sicurcantiere.db import Progetto
from sicurcantiere.repository import Repository
from tests.test_ai import ClientOpenAIFinto

FASE_ARCHIVIO = "Scavo di pulizia generale dell'area di cantiere"


def _parziale(**k):
    vuoto = {"cantiere": {}, "soggetti": [], "imprese": [], "fasi": [], "interferenze": [], "richieste_pos": []}
    vuoto.update(k)
    return vuoto


def _analisi():
    prima = _parziale(
        cantiere={"nome": "Scuola media", "comune": "Bari", "provincia": "ba", "importo_lavori": "€ 1.250.000,00",
                  "indirizzo": "Via Roma 1", "natura_opera": "", "cap": "", "titolo_abilitativo": "", "descrizione": ""},
        soggetti=[{"ruolo": "Coordinatore in fase di esecuzione", "nome": "Ing. Verdi", "qualifica": "",
                   "indirizzo": "", "telefono": "", "email": "verdi@pec.it"}],
        imprese=[{"ragione_sociale": "Alfa Costruzioni S.r.l.", "partita_iva": "01234567890", "ruolo": "Affidataria",
                  "sede": "", "datore_lavoro": "", "lavori_affidati": "Opere edili"}],
        fasi=[{"lavorazione": "Allestimento", "descrizione": FASE_ARCHIVIO, "impresa": "Alfa Costruzioni srl",
               "data_inizio": "2026-11-02", "data_fine": "2026-11-06", "zona": "A",
               "rischi": [{"rischio": "Presenza di sottoservizi", "misure": "Verificare le planimetrie"}],
               "prescrizioni": "Delimitare l'area di scavo"}],
        interferenze=[{"fasi": FASE_ARCHIVIO + " e montaggio gru", "imprese": "Alfa Costruzioni; Beta Gru",
                       "periodo": "novembre", "prescrizioni": "Sfasare le lavorazioni"}],
        richieste_pos=[{"impresa": "tutte", "richiesta": "Procedura di sollevamento dei carichi"}])
    seconda = _parziale(
        cantiere={"nome": "Altro nome", "cap": "70100"},
        imprese=[{"ragione_sociale": "ALFA COSTRUZIONI", "partita_iva": "", "ruolo": "", "sede": "Bari",
                  "datore_lavoro": "Mario Bianchi", "lavori_affidati": ""}],
        fasi=[{"lavorazione": "Impianti", "descrizione": "Posa delle canalizzazioni speciali", "impresa": "Beta Gru",
               "data_inizio": "", "data_fine": "", "zona": "",
               "rischi": [{"rischio": "Caduta di materiale dall'alto", "misure": "Casco"}], "prescrizioni": ""}],
        richieste_pos=[{"impresa": "Beta Gru", "richiesta": "Piano di sollevamento"}])
    return importa_psc.unisci([prima, seconda], file="PSC scuola.pdf")


def test_pulizia_e_parti():
    corpo = "\n".join(f"Riga {k} del corpo" for k in range(8))  # righe simili nel testo: restano
    pagine = [f"STUDIO TECNICO ROSSI - PSC rev. 2\nPagina {n}: testo con contenuto utile\n{corpo}\n"
              f"Fase {n} durata 10 giorni\nAltro testo\nUltima riga utile {n}\nPag. {n} di 10" for n in range(1, 11)]
    pagine[1] += "\n1. Premessa ........................ 3"
    pulito = importa_psc.pulisci_testo("\f".join(pagine))
    assert "STUDIO TECNICO" not in pulito and "Pag." not in pulito and "Premessa ...." not in pulito
    assert pulito.count("Riga 3 del corpo") == 10 and pulito.count("durata 10 giorni") == 10
    pezzi = importa_psc.parti("\n".join("x" * 100 for _ in range(1000)), 10_000)
    assert len(pezzi) == 11 and all(len(p) <= 10_000 for p in pezzi)
    assert importa_psc.stima_token("a" * 36_000) > 10_000


def test_testo_da_word(tmp_path):
    doc = Document()
    for n in range(40):
        doc.add_paragraph(f"Paragrafo {n} del piano di sicurezza e coordinamento del cantiere")
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Impresa", "Alfa Costruzioni"
    doc.save(tmp_path / "psc.docx")
    testo = importa_psc.estrai_testo(tmp_path / "psc.docx")
    assert "Impresa | Alfa Costruzioni" in testo


def test_unione_delle_parti():
    a = _analisi()
    assert a.cantiere["nome"] == "Scuola media" and a.cantiere["cap"] == "70100"
    assert [i["ragione_sociale"] for i in a.imprese] == ["Alfa Costruzioni S.r.l.", "Beta Gru"]
    assert a.imprese[0]["datore_lavoro"] == "Mario Bianchi"  # completata dalla seconda parte
    assert [f["descrizione"] for f in a.fasi_impresa("Alfa Costruzioni")] == [FASE_ARCHIVIO]


def test_lettura_con_assistente_in_una_volta():
    client = ClientOpenAIFinto(_parziale(imprese=[{"ragione_sociale": "Alfa", "partita_iva": "", "ruolo": "",
                                                   "sede": "", "datore_lavoro": "", "lavori_affidati": ""}]))
    testo = "\n".join("riga del PSC " * 8 for _ in range(600))
    analisi = AssistenteAI(fornitore=OPENAI, client=client).analizza_psc(testo)
    assert len(client.richieste) == len(importa_psc.parti(testo)) > 1
    assert [i["ragione_sociale"] for i in analisi.imprese] == ["Alfa"]


def test_file_del_pos(tmp_path):
    from sicurcantiere.documenti.pos import chiave_testo, genera_pos

    a = _analisi()
    file = importa_psc.crea_spo(a, a.imprese[0], tmp_path / "pos.spo")
    repo = Repository(Progetto(file))
    try:
        c = repo.cantiere()
        assert c["nome"] == "Scuola media" and c["provincia"] == "BA" and c["importo_lavori"] == 1_250_000
        cse = repo.elenco("soggetti")[0]
        assert cse["ruolo"] == "Coordinatore per l'esecuzione (CSE)" and cse["pec"] == "verdi@pec.it"
        impresa = repo.elenco("imprese")[0]
        assert impresa["tipo"] == "Affidataria" and impresa["datore_lavoro"] == "Mario Bianchi"
        fase = repo.elenco("fasi")[0]
        # riconosciuta nell'archivio: porta con sé macchine e rischi, più quelli scritti nel PSC
        assert fase["impresa_id"] == impresa["id"] and fase["data_inizio"] == "2026-11-02"
        assert repo.elenco("attrezzature", fase_id=fase["id"])
        assert "Presenza di sottoservizi" in [r["descrizione"] for r in repo.rischi(fase_id=fase["id"])]
        assert "Delimitare l'area di scavo" in fase["procedure_psc"] and "Sfasare" in fase["procedure_psc"]
        procedure = repo.testo(chiave_testo(impresa["id"], "procedure"))
        assert "sollevamento dei carichi" in procedure and "Piano di sollevamento" not in procedure
        assert repo.testo(importa_psc.CHIAVE_ORIGINE).startswith("Preparato dal PSC «PSC scuola.pdf»")
        assert "Beta Gru" in repo.testo(importa_psc.CHIAVE_ANALISI)
        genera_pos(repo, impresa["id"], tmp_path / "pos.docx")
    finally:
        repo.progetto.chiudi()
    # fase non in archivio, scelta a mano: rischi del PSC da valutare
    file = importa_psc.crea_spo(a, a.imprese[1], tmp_path / "beta.spo", fasi=a.fasi[1:])
    repo = Repository(Progetto(file))
    try:
        fase = repo.elenco("fasi")[0]
        assert fase["descrizione"] == "Posa delle canalizzazioni speciali"
        assert "verificare probabilità e danno" in fase["procedure_psc"]
    finally:
        repo.progetto.chiudi()


def test_parte_troncata_divisa_e_riletta():
    from types import SimpleNamespace

    risposte = iter(["troncata"] + [_parziale(imprese=[{"ragione_sociale": f"Impresa {n}", "partita_iva": "",
                                                        "ruolo": "", "sede": "", "datore_lavoro": "",
                                                        "lavori_affidati": ""}]) for n in range(10)])

    def create(**_):
        r = next(risposte)
        motivo = "length" if r == "troncata" else "stop"
        import json
        return SimpleNamespace(choices=[SimpleNamespace(finish_reason=motivo, message=SimpleNamespace(
            content="{}" if r == "troncata" else json.dumps(r), refusal=None))])

    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    testo = "\n".join("riga del PSC con molte fasi " * 4 for _ in range(100))  # una sola parte
    analisi = AssistenteAI(fornitore=OPENAI, client=client).analizza_psc(testo)
    assert [i["ragione_sociale"] for i in analisi.imprese] == ["Impresa 0", "Impresa 1"]
