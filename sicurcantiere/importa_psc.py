"""POS da un PSC fatto con qualsiasi programma (PDF o Word).

Il consulente che redige il POS di un'impresa riceve il PSC del coordinatore. Il programma:
1. estrae il testo dal PDF o dal Word e lo ripulisce in locale (intestazioni e piè di pagina
   ripetuti, numeri di pagina, indice, spazi), così all'assistente AI arriva solo il necessario;
2. lo fa leggere una volta sola all'assistente AI, a pezzi se è lungo, ricavando cantiere,
   soggetti, imprese, lavorazioni con rischi e prescrizioni, interferenze e procedure richieste
   ai POS: dalla stessa lettura si preparano i POS di tutte le imprese;
3. per ogni impresa scelta crea un file .spo da completare in SicurCantiere POS. Le lavorazioni
   riconosciute nell'archivio portano con sé macchine, attrezzi, rischi e misure (senza altre
   richieste all'AI); quelle non riconosciute riportano i rischi e le misure scritti nel PSC.

Il risultato della lettura resta nel file del POS (sezione «analisi_psc»), per poter verificare
in seguito che il POS risponda a ogni richiesta del PSC.
"""

from __future__ import annotations

import difflib
import json
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

from sicurcantiere import anagrafiche

CARATTERI_PER_PARTE = 12_000  # circa 3-4 mila token per richiesta: risposte che stanno nel limite
CARATTERI_PER_TOKEN = 3.6  # stima per l'italiano
CHIAVE_ANALISI = "analisi_psc"
CHIAVE_ORIGINE = "origine_psc"
SOGLIA_ARCHIVIO = 0.86  # somiglianza minima per riconoscere una fase dell'archivio


class ErroreImportazione(Exception):
    """Documento non leggibile o senza i dati che servono."""


# --- testo del documento ---

def estrai_testo(percorso: str | Path, avanzamento=None) -> str:
    """Testo del PDF o del Word; `avanzamento(pagina, pagine)` viene chiamata a ogni pagina del PDF."""
    percorso = Path(percorso)
    estensione = percorso.suffix.lower()
    if estensione == ".pdf":
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        try:
            lettore = PdfReader(str(percorso))
            pagine = []
            for n, pagina in enumerate(lettore.pages, 1):
                if avanzamento:
                    avanzamento(n, len(lettore.pages))
                pagine.append(pagina.extract_text() or "")
        except (PdfReadError, OSError, ValueError) as e:
            raise ErroreImportazione(f"Il PDF non si può leggere: {e}") from e
        testo = "\f".join(pagine)
    elif estensione == ".docx":
        from docx import Document

        try:
            doc = Document(str(percorso))
        except Exception as e:  # noqa: BLE001 - file Word danneggiato o non valido
            raise ErroreImportazione(f"Il file Word non si può leggere: {e}") from e
        righe = [p.text for p in doc.paragraphs]
        for tabella in doc.tables:
            for riga in tabella.rows:
                celle = list(dict.fromkeys(c.text.strip() for c in riga.cells if c.text.strip()))
                righe.append(" | ".join(celle))
        testo = "\n".join(righe)
    elif estensione in (".txt", ".md"):
        testo = percorso.read_text(encoding="utf-8", errors="replace")
    else:
        raise ErroreImportazione("Formato non supportato: servono un PDF o un file Word (.docx).")
    if len(re.sub(r"\s", "", testo)) < 500:
        raise ErroreImportazione("Nel documento non c'è testo da leggere: probabilmente è una scansione. "
                                 "Serve il PDF originale o il file Word del PSC.")
    return testo


def pulisci_testo(testo: str) -> str:
    """Toglie in locale ciò che non serve all'analisi, per ridurre i token inviati."""
    pagine = testo.split("\f")
    righe_pagine = [[r.strip() for r in p.splitlines()] for p in pagine]
    # intestazioni e piè di pagina: tra le prime e le ultime righe di ogni pagina, quelle uguali (a meno
    # dei numeri) su molte pagine; il corpo del testo non si tocca
    def bordi(righe: list[str]) -> list[str]:
        piene = [r for r in righe if r]
        return piene[:3] + piene[-3:]

    ripetute: set[str] = set()
    if len(pagine) >= 4:
        conta = Counter(re.sub(r"\d+", "#", r) for righe in righe_pagine for r in set(bordi(righe)))
        ripetute = {r for r, n in conta.items() if n >= max(3, len(pagine) // 3) and len(r) < 160}
    uscita = []
    for righe in righe_pagine:
        margini = set(bordi(righe))
        for r in righe:
            if not r or (r in margini and re.sub(r"\d+", "#", r) in ripetute):
                continue
            if re.fullmatch(r"(pag(ina)?\.?\s*)?\d+(\s*(di|/)\s*\d+)?", r, re.I):
                continue
            if re.search(r"\.{5,}\s*\d+$", r):  # righe dell'indice «Titolo ........ 12»
                continue
            uscita.append(re.sub(r"[ \t]{2,}", " ", r))
    return "\n".join(uscita)


def parti(testo: str, dimensione: int = CARATTERI_PER_PARTE) -> list[str]:
    """Divide il testo in parti di dimensione massima, senza spezzare le righe."""
    risultato, corrente, lunghezza = [], [], 0
    for riga in testo.splitlines():
        if lunghezza + len(riga) > dimensione and corrente:
            risultato.append("\n".join(corrente))
            corrente, lunghezza = [], 0
        corrente.append(riga)
        lunghezza += len(riga) + 1
    if corrente:
        risultato.append("\n".join(corrente))
    return risultato


def stima_token(testo: str) -> int:
    """Token inviati più quelli della risposta, all'incirca."""
    ingresso = len(testo) / CARATTERI_PER_TOKEN
    return int(ingresso * 1.15 + 1500 * len(parti(testo)))


# --- schema della lettura ---

def _stringa(descrizione: str) -> dict[str, Any]:
    return {"type": "string", "description": descrizione}


def _oggetto(proprieta: dict[str, Any]) -> dict[str, Any]:
    return {"type": "object", "properties": proprieta, "required": list(proprieta), "additionalProperties": False}


def _elenco(proprieta: dict[str, Any], descrizione: str) -> dict[str, Any]:
    return {"type": "array", "description": descrizione, "items": _oggetto(proprieta)}


SCHEMA = _oggetto({
    "cantiere": _oggetto({
        "nome": _stringa("Denominazione dell'opera o oggetto dei lavori"),
        "natura_opera": _stringa("Natura dell'opera (es. nuova costruzione, ristrutturazione)"),
        "indirizzo": _stringa("Via e numero civico del cantiere"),
        "comune": _stringa("Comune del cantiere"),
        "provincia": _stringa("Sigla della provincia, due lettere"),
        "cap": _stringa("CAP"),
        "importo_lavori": _stringa("Importo presunto dei lavori in euro, solo il numero"),
        "titolo_abilitativo": _stringa("Permesso di costruire, SCIA, CILA... con numero e data"),
        "descrizione": _stringa("Descrizione sintetica dell'opera, al massimo 5 righe"),
    }),
    "soggetti": _elenco({
        "ruolo": _stringa("Committente, Responsabile dei lavori, CSP, CSE, Direttore dei lavori, Progettista..."),
        "nome": _stringa("Nome o ragione sociale"),
        "qualifica": _stringa("Titolo professionale, se indicato"),
        "indirizzo": _stringa("Indirizzo, se indicato"),
        "telefono": _stringa("Telefono, se indicato"),
        "email": _stringa("Email o PEC, se indicata"),
    }, "Soggetti con compiti di sicurezza indicati nel PSC"),
    "imprese": _elenco({
        "ragione_sociale": _stringa("Ragione sociale dell'impresa o nome del lavoratore autonomo"),
        "partita_iva": _stringa("Partita IVA, se indicata"),
        "ruolo": _stringa("Affidataria, Esecutrice, Subappaltatrice, Lavoratore autonomo o vuoto"),
        "sede": _stringa("Sede, se indicata"),
        "datore_lavoro": _stringa("Datore di lavoro o legale rappresentante, se indicato"),
        "lavori_affidati": _stringa("Lavori o categorie affidati all'impresa"),
    }, "Imprese e lavoratori autonomi nominati nel PSC (anche solo come «impresa da individuare»)"),
    "fasi": _elenco({
        "lavorazione": _stringa("Lavorazione o gruppo a cui appartiene la fase"),
        "descrizione": _stringa("Nome della fase di lavoro come scritto nel PSC"),
        "impresa": _stringa("Impresa che la esegue, come scritta nel PSC; vuoto se non indicata"),
        "data_inizio": _stringa("Inizio previsto in formato AAAA-MM-GG, vuoto se non indicato"),
        "data_fine": _stringa("Fine prevista in formato AAAA-MM-GG, vuoto se non indicata"),
        "zona": _stringa("Zona o area del cantiere, se indicata"),
        "rischi": _elenco({
            "rischio": _stringa("Rischio, in breve"),
            "misure": _stringa("Misure di prevenzione e protezione del PSC per questo rischio, al massimo 150 caratteri"),
        }, "Rischi della fase con le misure del PSC"),
        "prescrizioni": _stringa("Prescrizioni operative e DPI richiesti dal PSC per la fase, al massimo 200 caratteri"),
    }, "Fasi di lavoro del PSC o del cronoprogramma"),
    "interferenze": _elenco({
        "fasi": _stringa("Fasi o lavorazioni che interferiscono"),
        "imprese": _stringa("Imprese coinvolte, separate da punto e virgola; vuoto se non indicate"),
        "periodo": _stringa("Periodo, se indicato"),
        "prescrizioni": _stringa("Misure di coordinamento prescritte dal CSE"),
    }, "Interferenze tra lavorazioni e relative prescrizioni"),
    "richieste_pos": _elenco({
        "impresa": _stringa("Impresa a cui è rivolta la richiesta, oppure «tutte»"),
        "richiesta": _stringa("Procedura complementare, contenuto o documento che il POS deve riportare"),
    }, "Procedure complementari e di dettaglio, contenuti e documenti che il PSC chiede ai POS delle imprese"),
})

ISTRUZIONI = (
    "Leggi questa parte ({n} di {totale}) del piano di sicurezza e coordinamento (PSC) di un cantiere e "
    "riporta i dati richiesti dallo schema, così come sono scritti nel documento, in forma breve e senza "
    "ragionamenti o commenti: solo il JSON. Non inventare: se un dato "
    "non c'è lascia il campo vuoto o l'elenco vuoto. Riassumi i testi lunghi in modo fedele. Le altre parti "
    "del documento vengono lette a parte: riporta solo quello che trovi qui.\n\n<psc>\n{testo}\n</psc>"
)


# --- risultato ---

@dataclass
class Analisi:
    cantiere: dict[str, str] = field(default_factory=dict)
    soggetti: list[dict[str, str]] = field(default_factory=list)
    imprese: list[dict[str, str]] = field(default_factory=list)
    fasi: list[dict[str, Any]] = field(default_factory=list)
    interferenze: list[dict[str, str]] = field(default_factory=list)
    richieste_pos: list[dict[str, str]] = field(default_factory=list)
    file: str = ""
    saltati: list[str] = field(default_factory=list)  # inizi dei pezzi che l'assistente non è riuscito a leggere

    def json(self) -> str:
        return json.dumps(self.__dict__, ensure_ascii=False, indent=1)

    def fasi_impresa(self, impresa: str) -> list[dict[str, Any]]:
        return [f for f in self.fasi if stessa_impresa(f.get("impresa", ""), impresa)]


def _norma(nome: str) -> str:
    nome = nome.lower()
    nome = re.sub(r"\b(s\.?\s?r\.?\s?l\.?s?|s\.?\s?p\.?\s?a\.?|s\.?\s?n\.?\s?c\.?|s\.?\s?a\.?\s?s\.?|"
                  r"soc\.?|societ[aà]|coop(erativa)?|impresa|ditta|edile|costruzioni)\b", " ", nome)
    return " ".join(re.sub(r"[^\w]", " ", nome).split())


def stessa_impresa(a: str, b: str) -> bool:
    na, nb = _norma(a or ""), _norma(b or "")
    if not na or not nb:
        return False
    return na == nb or na in nb or nb in na or difflib.SequenceMatcher(None, na, nb).ratio() >= 0.85


def _testo(v: Any) -> str:
    return " ".join(str(v or "").split()) if not isinstance(v, str) or "\n" not in v else str(v).strip()


def unisci(parziali: list[dict[str, Any]], file: str = "") -> Analisi:
    """Unisce le letture delle parti del documento togliendo i doppioni."""
    a = Analisi(file=file)
    for p in parziali:
        for k, v in (p.get("cantiere") or {}).items():
            if _testo(v) and not a.cantiere.get(k):
                a.cantiere[k] = _testo(v)
        for s in p.get("soggetti") or []:
            s = {k: _testo(v) for k, v in s.items()}
            if s.get("nome") and not any(x["nome"].lower() == s["nome"].lower() and x["ruolo"] == s["ruolo"]
                                         for x in a.soggetti):
                a.soggetti.append(s)
        for i in p.get("imprese") or []:
            i = {k: _testo(v) for k, v in i.items()}
            if not i.get("ragione_sociale"):
                continue
            uguale = next((x for x in a.imprese if stessa_impresa(x["ragione_sociale"], i["ragione_sociale"])
                           or (i.get("partita_iva") and x.get("partita_iva") == i["partita_iva"])), None)
            if uguale is None:
                a.imprese.append(i)
            else:
                for k, v in i.items():
                    if v and not uguale.get(k):
                        uguale[k] = v
        for f in p.get("fasi") or []:
            if not _testo(f.get("descrizione")):
                continue
            chiave = (_testo(f["descrizione"]).lower(), _norma(_testo(f.get("impresa"))))
            uguale = next((x for x in a.fasi if (x["descrizione"].lower(), _norma(x["impresa"])) == chiave), None)
            rischi = [{"rischio": _testo(r.get("rischio")), "misure": _testo(r.get("misure"))}
                      for r in f.get("rischi") or [] if _testo(r.get("rischio"))]
            if uguale is None:
                a.fasi.append({**{k: _testo(v) for k, v in f.items() if k != "rischi"}, "rischi": rischi})
            else:
                noti = {r["rischio"].lower() for r in uguale["rischi"]}
                uguale["rischi"] += [r for r in rischi if r["rischio"].lower() not in noti]
                for k, v in f.items():
                    if k != "rischi" and _testo(v) and not uguale.get(k):
                        uguale[k] = _testo(v)
        for elenco, chiave in ((a.interferenze, "interferenze"), (a.richieste_pos, "richieste_pos")):
            for x in p.get(chiave) or []:
                x = {k: _testo(v) for k, v in x.items()}
                if any(x.values()) and x not in elenco:
                    elenco.append(x)
    # imprese nominate solo nelle fasi
    for f in a.fasi:
        if f["impresa"] and not any(stessa_impresa(i["ragione_sociale"], f["impresa"]) for i in a.imprese):
            if not re.search(r"da (individuare|definire|nominare)|n\.?d\.?|vari", f["impresa"], re.I):
                a.imprese.append({"ragione_sociale": f["impresa"], "partita_iva": "", "ruolo": "", "sede": "",
                                  "datore_lavoro": "", "lavori_affidati": ""})
    return a


# --- file del POS ---

def _ruolo_soggetto(ruolo: str) -> str:
    r = ruolo.lower()
    for parole, nome in ((("esecuzione", "cse"), anagrafiche.CSE), (("progettazione", "csp"), anagrafiche.CSP),
                         (("responsabile dei lavori", "responsabile lavori"), anagrafiche.RESPONSABILE_LAVORI),
                         (("committente",), anagrafiche.COMMITTENTE),
                         (("direttore dei lavori", "direttore lavori", "d.l."), "Direttore dei lavori")):
        if any(p in r for p in parole):
            return nome
    return ruolo.strip() or "Soggetto"


def _ruolo_impresa(ruolo: str) -> str:
    r = ruolo.lower()
    if "autonom" in r:
        return anagrafiche.AUTONOMO
    if "affidataria" in r and "esecut" in r:
        return "Affidataria ed esecutrice"
    for nome in ("Affidataria", "Subappaltatrice", "Esecutrice"):
        if nome.lower()[:8] in r:
            return nome
    return "Esecutrice"


def _data(valore: str) -> str | None:
    try:
        return date.fromisoformat(valore.strip()[:10]).isoformat() if valore else None
    except ValueError:
        return None


def _importo(valore: str) -> float:
    numero = re.sub(r"[^\d,\.]", "", valore or "")
    if "," in numero:
        numero = numero.replace(".", "").replace(",", ".")
    try:
        return float(numero) if numero else 0.0
    except ValueError:
        return 0.0


def fase_archivio(descrizione: str) -> dict[str, Any] | None:
    """La fase dell'archivio con lo stesso nome (o quasi), se c'è."""
    from sicurcantiere.archivio import LAVORAZIONI

    cercata = " ".join(descrizione.lower().split())
    migliore, punteggio = None, 0.0
    for l in LAVORAZIONI:
        r = difflib.SequenceMatcher(None, cercata, l["descrizione"].lower()).ratio()
        if r > punteggio:
            migliore, punteggio = l, r
    return migliore if punteggio >= SOGLIA_ARCHIVIO else None


def _riguarda(voce: dict[str, str], impresa: str, nomi_fasi: list[str]) -> bool:
    imprese = [x.strip() for x in re.split(r";|,| e ", voce.get("imprese", "") or voce.get("impresa", "")) if x.strip()]
    if any(x.lower() in ("tutte", "tutte le imprese") for x in imprese):
        return True
    if any(stessa_impresa(x, impresa) for x in imprese):
        return True
    testo = (voce.get("fasi", "") or "").lower()
    return bool(testo) and any(n.lower() in testo for n in nomi_fasi if len(n) > 6)


def crea_spo(analisi: Analisi, impresa: dict[str, str], destinazione: str | Path,
             fasi: list[dict[str, Any]] | None = None) -> Path:
    """File .spo del POS dell'impresa con i dati letti dal PSC.

    `fasi` sono le fasi del PSC svolte dall'impresa; se non indicate, quelle che il PSC le assegna.
    """
    from sicurcantiere.db import Progetto
    from sicurcantiere.documenti.comune import data_it
    from sicurcantiere.documenti.pos import chiave_testo
    from sicurcantiere.repository import Repository

    destinazione = Path(destinazione)
    progetto = Progetto(destinazione, crea=True)
    try:
        repo = Repository(progetto)
        c = analisi.cantiere
        repo.aggiorna_cantiere(
            nome=c.get("nome", ""), natura_opera=c.get("natura_opera", ""), indirizzo=c.get("indirizzo", ""),
            comune=c.get("comune", ""), provincia=c.get("provincia", "")[:2].upper(), cap=c.get("cap", ""),
            descrizione=c.get("descrizione", ""), titolo_abilitativo=c.get("titolo_abilitativo", ""),
            importo_lavori=_importo(c.get("importo_lavori", "")))
        for s in analisi.soggetti:
            email = s.get("email", "")
            repo.inserisci("soggetti", ruolo=_ruolo_soggetto(s.get("ruolo", "")), nome=s.get("nome", ""),
                           qualifica=s.get("qualifica", ""), indirizzo=s.get("indirizzo", ""),
                           telefono=s.get("telefono", ""), email="" if "pec" in email.lower() else email,
                           pec=email if "pec" in email.lower() else "")
        id_impresa = repo.inserisci(
            "imprese", ragione_sociale=impresa["ragione_sociale"], partita_iva=impresa.get("partita_iva", ""),
            tipo=_ruolo_impresa(impresa.get("ruolo", "")), indirizzo=impresa.get("sede", ""),
            datore_lavoro=impresa.get("datore_lavoro", ""), tipologia_lavori=impresa.get("lavori_affidati", ""))

        fasi = analisi.fasi_impresa(impresa["ragione_sociale"]) if fasi is None else fasi
        nomi_fasi = [f["descrizione"] for f in fasi]
        interferenze = [i for i in analisi.interferenze if _riguarda(i, impresa["ragione_sociale"], nomi_fasi)]
        for f in fasi:
            modello = fase_archivio(f["descrizione"])
            if modello is not None:
                fase_id = repo.importa_lavorazione(modello)
            else:
                fase_id = repo.inserisci("fasi", codice=repo.codice_automatico(f.get("lavorazione", "")),
                                         descrizione=f["descrizione"],
                                         lavorazione=f.get("lavorazione", ""))
            noti = {r["descrizione"].lower() for r in repo.rischi(fase_id=fase_id)}
            for r in f["rischi"]:
                if r["rischio"].lower() not in noti:
                    repo.inserisci("rischi", fase_id=fase_id, descrizione=r["rischio"], misure=r["misure"],
                                   probabilita=2, danno=2)
            procedure = []
            if f.get("prescrizioni"):
                procedure.append(f"Prescrizioni del PSC: {f['prescrizioni']}")
            for i in interferenze:
                if not i.get("fasi") or f["descrizione"].lower() in i["fasi"].lower() or \
                        (f.get("lavorazione") and f["lavorazione"].lower() in i["fasi"].lower()):
                    procedure.append(f"Interferenza ({i['fasi'] or 'vedi PSC'}"
                                     + (f", {i['periodo']}" if i.get("periodo") else "") + f"): {i['prescrizioni']}")
            if f["rischi"] and modello is None:
                procedure.append("Rischi riportati dal PSC: verificare probabilità e danno nella valutazione.")
            repo.aggiorna("fasi", fase_id, impresa_id=id_impresa, data_inizio=_data(f.get("data_inizio", "")),
                          data_fine=_data(f.get("data_fine", "")), zona=f.get("zona", ""),
                          procedure_psc="\n".join(procedure),
                          **({"lavorazione": f["lavorazione"]} if f.get("lavorazione") else {}))
        richieste = [r["richiesta"] for r in analisi.richieste_pos
                     if r.get("richiesta") and _riguarda({"imprese": r.get("impresa", "")},
                                                         impresa["ragione_sociale"], [])]
        if richieste:
            repo.salva_testo(chiave_testo(id_impresa, "procedure"),
                             "Procedure complementari e di dettaglio richieste dal PSC:\n"
                             + "\n".join(f"- {r}" for r in richieste))
        nome_file = Path(analisi.file).name if analisi.file else "PSC"
        repo.salva_testo(CHIAVE_ORIGINE, f"Preparato dal PSC «{nome_file}» il {data_it(date.today())}")
        repo.salva_testo(CHIAVE_ANALISI, analisi.json())
    finally:
        progetto.chiudi()
    return destinazione
