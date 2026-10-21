"""Impaginazione dei documenti: stili, copertina, capitoli numerati, schede di dati, tabelle, immagini.

Impostazione grafica propria di SicurCantiere, con gli stessi colori del programma:
- capitoli numerati (1, 1.1) allineati a sinistra con un filetto nel colore del programma e il
  riferimento di legge sotto, in piccolo; le parti del documento si aprono su una pagina nuova;
- dati in schede etichetta-valore a righe sottili, tabelle leggere senza griglia verticale con
  l'intestazione ripetuta a ogni pagina;
- lavorazioni e fasi in riquadri con barra laterale, note e avvertenze in box evidenziati;
- intestazione con documento e revisione, piè di pagina con cantiere e «Pagina X di Y».
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Caratteri Ubuntu (Ubuntu Font Licence): l'installatore li installa in Windows, così Word li mostra;
# salvando da Word con «incorpora i caratteri» li vede anche chi riceve il documento.
CARATTERE = "Ubuntu"
CARATTERE_TITOLI = "Ubuntu Medium"
LARGHEZZA_UTILE = 17.4  # cm, A4 con margini di 1,8 cm

# Colori (gli stessi del programma POS); ACCENTO e TINTA si possono cambiare per ogni documento.
ACCENTO, ACCENTO_SCURO, TINTA = "1C6597", "134A70", "E6F1F9"
INCHIOSTRO, TESTO2, TESTO3 = "17202A", "45515E", "6B7785"
LINEA, LINEA_FORTE, SUPERFICIE = "D7DDE3", "B9C3CD", "F4F6F8"
AVVISO, AVVISO_SFONDO = "9A4500", "FFF2DC"
# Livelli di rischio: (sfondo, testo) come nel programma.
LIVELLI = [("DCF0E0", "1B5E2B"), ("FDEFB2", "674B00"), ("FFDCC0", "863300"), ("F7D2CF", "8F1A13")]


def rgb(colore: str) -> RGBColor:
    return RGBColor.from_string(colore)


# Nomi usati dal codice precedente.
BLU = rgb(ACCENTO)
GRIGIO = rgb(TESTO3)


def imposta_colore(accento: str, scuro: str, tinta: str) -> None:
    """Colore del programma per il documento (POS blu, PSC e PiMUS con i loro colori)."""
    global ACCENTO, ACCENTO_SCURO, TINTA, BLU
    ACCENTO, ACCENTO_SCURO, TINTA = accento, scuro, tinta
    BLU = rgb(accento)


# --- utilità XML ---

def _ombra(elemento_pr, colore: str) -> None:
    sfondo = OxmlElement("w:shd")
    sfondo.set(qn("w:val"), "clear")
    sfondo.set(qn("w:color"), "auto")
    sfondo.set(qn("w:fill"), colore)
    elemento_pr.append(sfondo)


def _bordo_paragrafo(paragrafo, lati: tuple[str, ...], colore: str, spessore: int = 8, spazio: int = 4) -> None:
    pr = paragrafo._p.get_or_add_pPr()
    bordi = OxmlElement("w:pBdr")
    for lato in lati:
        b = OxmlElement(f"w:{lato}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(spessore))
        b.set(qn("w:space"), str(spazio))
        b.set(qn("w:color"), colore)
        bordi.append(b)
    pr.append(bordi)


def ombreggia_paragrafo(paragrafo, colore: str) -> None:
    _ombra(paragrafo._p.get_or_add_pPr(), colore)


def ombreggia_cella(cella, colore: str) -> None:
    _ombra(cella._tc.get_or_add_tcPr(), colore)


def larghezze_colonne(t, larghezze_cm: list[float]) -> None:
    """Fissa la larghezza delle colonne (griglia e celle), rispettata anche da LibreOffice."""
    t.autofit = False
    griglia = t._tbl.tblGrid
    for col, cm in zip(griglia.findall(qn("w:gridCol")), larghezze_cm):
        col.set(qn("w:w"), str(int(Cm(cm).twips)))
    for riga in t.rows:
        for cella, cm in zip(riga.cells, larghezze_cm):
            cella.width = Cm(cm)


def _bordi_tabella(tabella, **lati: tuple[str, int] | None) -> None:
    """Bordi della tabella: per ogni lato (top, left, bottom, right, insideH, insideV) colore e spessore."""
    pr = tabella._tbl.tblPr
    for vecchio in pr.findall(qn("w:tblBorders")):
        pr.remove(vecchio)
    bordi = OxmlElement("w:tblBorders")
    for lato in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{lato}")
        valore = lati.get(lato)
        if valore is None:
            b.set(qn("w:val"), "nil")
        else:
            b.set(qn("w:val"), "single")
            b.set(qn("w:color"), valore[0])
            b.set(qn("w:sz"), str(valore[1]))
            b.set(qn("w:space"), "0")
        bordi.append(b)
    pr.append(bordi)


def senza_bordi(tabella) -> None:
    _bordi_tabella(tabella)


def _margini_celle(tabella, sopra: float = 0.08, lati: float = 0.15) -> None:
    pr = tabella._tbl.tblPr
    margini = OxmlElement("w:tblCellMar")
    for lato, cm in (("top", sopra), ("bottom", sopra), ("left", lati), ("right", lati)):
        m = OxmlElement(f"w:{lato}")
        m.set(qn("w:w"), str(int(Cm(cm).twips)))
        m.set(qn("w:type"), "dxa")
        margini.append(m)
    pr.append(margini)


def bordi_cella(cella, colore: str = "000000", spessore: int = 12, lati=("top", "left", "bottom", "right")) -> None:
    pr = cella._tc.get_or_add_tcPr()
    bordi = OxmlElement("w:tcBorders")
    for lato in lati:
        b = OxmlElement(f"w:{lato}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(spessore))
        b.set(qn("w:color"), colore)
        bordi.append(b)
    pr.append(bordi)


def _ripeti_intestazione(riga) -> None:
    pr = riga._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    pr.append(e)


def _non_spezzare(riga) -> None:
    pr = riga._tr.get_or_add_trPr()
    pr.append(OxmlElement("w:cantSplit"))


def campo(paragrafo, istruzione: str, testo: str = ""):
    """Campo di Word (numero di pagina, indice...) aggiornato all'apertura del documento."""
    run = paragrafo.add_run()
    inizio = OxmlElement("w:fldChar")
    inizio.set(qn("w:fldCharType"), "begin")
    istr = OxmlElement("w:instrText")
    istr.set(qn("xml:space"), "preserve")
    istr.text = istruzione
    separa = OxmlElement("w:fldChar")
    separa.set(qn("w:fldCharType"), "separate")
    segnaposto = OxmlElement("w:t")
    segnaposto.text = testo
    fine = OxmlElement("w:fldChar")
    fine.set(qn("w:fldCharType"), "end")
    for e in (inizio, istr, separa, segnaposto, fine):
        run._r.append(e)
    return run


def incorpora_caratteri(doc) -> None:
    """Ubuntu come carattere di base del documento (anche dove gli stili non lo indicano) e richiesta a
    Word di incorporare i caratteri al salvataggio, così chi riceve il file vede il font Ubuntu."""
    predefiniti = doc.styles.element.find(qn("w:docDefaults"))
    if predefiniti is not None:
        rpr = predefiniti.find(qn("w:rPrDefault"))
        rpr = rpr.find(qn("w:rPr")) if rpr is not None else None
        if rpr is not None:
            for vecchio in rpr.findall(qn("w:rFonts")):
                rpr.remove(vecchio)
            caratteri = OxmlElement("w:rFonts")
            for attributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                caratteri.set(qn(attributo), CARATTERE)
            rpr.insert(0, caratteri)
    # i caratteri «del tema» (Calibri) negli stili prevalgono in Word su quelli indicati: si tolgono
    for caratteri in doc.styles.element.iter(qn("w:rFonts")):
        tema = [k for k in caratteri.attrib if k.endswith("Theme") or k.endswith("theme")]
        for k in tema:
            del caratteri.attrib[k]
        if tema and caratteri.get(qn("w:ascii")) is None:
            caratteri.set(qn("w:ascii"), CARATTERE)
            caratteri.set(qn("w:hAnsi"), CARATTERE)
    impostazioni = doc.settings.element
    for nome in ("w:saveSubsetFonts", "w:embedTrueTypeFonts"):
        if impostazioni.find(qn(nome)) is None:
            impostazioni.insert(0, OxmlElement(nome))


def aggiorna_campi_all_apertura(doc) -> None:
    impostazioni = doc.settings.element
    elemento = OxmlElement("w:updateFields")
    elemento.set(qn("w:val"), "true")
    impostazioni.append(elemento)


def _run(paragrafo, testo: str, dimensione: float | None = None, colore: str | None = None,
         grassetto: bool | None = None, corsivo: bool | None = None, carattere: str | None = None):
    r = paragrafo.add_run(testo)
    if dimensione:
        r.font.size = Pt(dimensione)
    if colore:
        r.font.color.rgb = rgb(colore)
    if grassetto is not None:
        r.bold = grassetto
    if corsivo is not None:
        r.italic = corsivo
    if carattere:
        r.font.name = carattere
    return r


def _spaziatura(paragrafo, prima: float = 0, dopo: float = 0) -> None:
    paragrafo.paragraph_format.space_before = Pt(prima)
    paragrafo.paragraph_format.space_after = Pt(dopo)


# --- stili del documento ---

def imposta_stili(doc) -> None:
    incorpora_caratteri(doc)
    normale = doc.styles["Normal"]
    normale.font.name = CARATTERE
    normale.font.size = Pt(9.5)
    normale.font.color.rgb = rgb(INCHIOSTRO)
    normale.paragraph_format.space_after = Pt(4)
    normale.paragraph_format.line_spacing = 1.12

    for nome, dimensione, colore, prima, dopo in (("Heading 1", 17, INCHIOSTRO, 20, 2),
                                                   ("Heading 2", 12.5, ACCENTO_SCURO, 14, 4),
                                                   ("Heading 3", 10.5, INCHIOSTRO, 10, 3)):
        h = doc.styles[nome]
        h.font.name, h.font.size, h.font.bold, h.font.italic = CARATTERE_TITOLI, Pt(dimensione), False, False
        h.font.color.rgb = rgb(colore)
        h.font.all_caps = False
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h.paragraph_format.space_before = Pt(prima)
        h.paragraph_format.space_after = Pt(dopo)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.page_break_before = False
    for nome in ("List Bullet", "List Number"):
        doc.styles[nome].font.name = CARATTERE
        doc.styles[nome].paragraph_format.space_after = Pt(2)


def imposta_pagina(doc, piede: str, intestazione: tuple[str, str] = ("", "")) -> None:
    """A4 con intestazione (documento a sinistra, revisione a destra) e piè di pagina con
    cantiere e «Pagina X di Y». La copertina resta senza intestazione e piè di pagina."""
    for sezione in doc.sections:
        sezione.page_height, sezione.page_width = Cm(29.7), Cm(21)
        sezione.left_margin = sezione.right_margin = Cm(1.8)
        sezione.top_margin, sezione.bottom_margin = Cm(2.0), Cm(1.8)
        sezione.header_distance, sezione.footer_distance = Cm(0.9), Cm(0.8)
        sezione.different_first_page_header_footer = True
        _intestazione(sezione.header.paragraphs[0], *intestazione)
        _piede(sezione.footer.paragraphs[0], piede)


def _tabulazione_destra(paragrafo) -> None:
    pr = paragrafo._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    for posizione in (4680, 9360, 4513, 9026):  # tabulazioni degli stili Intestazione e Piè di pagina
        via = OxmlElement("w:tab")
        via.set(qn("w:val"), "clear")
        via.set(qn("w:pos"), str(posizione))
        tabs.append(via)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(int(Cm(LARGHEZZA_UTILE).twips)))
    tabs.append(tab)
    pr.append(tabs)


def _intestazione(p, sinistra: str, destra: str) -> None:
    p.text = ""
    _tabulazione_destra(p)
    _bordo_paragrafo(p, ("bottom",), LINEA, 4, 3)
    _run(p, sinistra, 7.5, TESTO3, grassetto=True)
    _run(p, "\t" + destra, 7.5, TESTO3)


def _piede(p, testo: str) -> None:
    p.text = ""
    _tabulazione_destra(p)
    _bordo_paragrafo(p, ("top",), LINEA, 4, 3)
    _run(p, testo, 7.5, TESTO3)
    _run(p, "\tPagina ", 7.5, TESTO3)
    campo(p, "PAGE", "1").font.size = Pt(7.5)
    _run(p, " di ", 7.5, TESTO3)
    campo(p, "NUMPAGES", "1").font.size = Pt(7.5)


# --- numerazione dei capitoli ---

def _numeri(doc) -> list[int]:
    if not hasattr(doc, "_numeri_sc"):
        doc._numeri_sc = [0, 0, 0]
    return doc._numeri_sc


def parte(doc, numero: str, titolo: str, descrizione: str = "") -> None:
    """Apertura di una parte del documento su pagina nuova (non entra nell'indice)."""
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    _spaziatura(p, 0, 2)
    _run(p, f"PARTE {numero}", 8.5, ACCENTO, grassetto=True)
    p = doc.add_paragraph()
    _spaziatura(p, 0, 4 if descrizione else 14)
    _bordo_paragrafo(p, ("bottom",), ACCENTO, 18, 6)
    _run(p, titolo, 22, INCHIOSTRO, carattere=CARATTERE_TITOLI)
    if descrizione:
        p = doc.add_paragraph()
        _spaziatura(p, 0, 14)
        _run(p, descrizione, 9, TESTO2)


def capitolo(doc, titolo: str, riferimento: str = "", nuova_pagina: bool = False) -> None:
    """Capitolo numerato (1, 2, 3...) con filetto colorato e riferimento di legge sotto."""
    n = _numeri(doc)
    n[0], n[1], n[2] = n[0] + 1, 0, 0
    h = doc.add_heading(level=1)
    h.paragraph_format.page_break_before = nuova_pagina
    _bordo_paragrafo(h, ("bottom",), ACCENTO, 12, 4)
    _run(h, f"{n[0]}", colore=ACCENTO)
    _run(h, f"   {titolo}")
    if riferimento:
        p = doc.add_paragraph()
        _spaziatura(p, 2, 10)
        p.paragraph_format.keep_with_next = True
        _run(p, riferimento[0].upper() + riferimento[1:], 8, TESTO3)
    else:
        h.paragraph_format.space_after = Pt(10)


def sezione(doc, titolo: str, sottotitolo: str = "") -> None:
    """Paragrafo numerato (1.1, 1.2...) dentro il capitolo, incluso nell'indice."""
    n = _numeri(doc)
    n[1], n[2] = n[1] + 1, 0
    h = doc.add_heading(level=2)
    _run(h, f"{n[0]}.{n[1]}  ", colore=TESTO3)
    _run(h, titolo)
    if sottotitolo:
        _run(h, f"  {sottotitolo}", 9, TESTO3, carattere=CARATTERE)


def sottosezione(doc, titolo: str) -> None:
    """Titolo di terzo livello (non numerato, fuori dall'indice)."""
    doc.add_heading(titolo, 3)


def banda(doc, testo: str) -> None:
    """Sottotitolo di gruppo: maiuscoletto nel colore del programma con filetto sottile."""
    p = doc.add_paragraph()
    _spaziatura(p, 10, 4)
    p.paragraph_format.keep_with_next = True
    _bordo_paragrafo(p, ("bottom",), LINEA, 4, 2)
    r = _run(p, testo, 8, ACCENTO_SCURO, grassetto=True)
    r.font.all_caps = True
    spaziatura = OxmlElement("w:spacing")
    spaziatura.set(qn("w:val"), "10")
    r._r.get_or_add_rPr().append(spaziatura)


def scheda(doc, voci: list[tuple[str, Any]], larghezza_etichetta: float = 5.2, titolo: str = "") -> None:
    """Elenco etichetta-valore a righe sottili; le voci vuote non compaiono, e se sono tutte vuote
    non compare neanche il titolo del gruppo."""
    voci = [(e, str(v).strip()) for e, v in voci if v not in (None, "", "-") and str(v).strip()]
    if not voci:
        return
    if titolo:
        banda(doc, titolo)
    t = doc.add_table(rows=0, cols=2)
    _bordi_tabella(t, insideH=(LINEA, 4), bottom=(LINEA, 4))
    _margini_celle(t, 0.06, 0.1)
    for etichetta, valore in voci:
        riga = t.add_row()
        _non_spezzare(riga)
        celle = riga.cells
        p = celle[0].paragraphs[0]
        _spaziatura(p)
        _run(p, etichetta, 8.5, TESTO3)
        for n, testo in enumerate(valore.split("\n")):
            p = celle[1].paragraphs[0] if n == 0 else celle[1].add_paragraph()
            _spaziatura(p)
            _run(p, testo, 9.5, INCHIOSTRO)
    larghezze_colonne(t, [larghezza_etichetta, LARGHEZZA_UTILE - larghezza_etichetta])
    _spazio_dopo(doc)


def _spazio_dopo(doc, punti: float = 4) -> None:
    p = doc.add_paragraph()
    _spaziatura(p)
    p.paragraph_format.line_spacing = Pt(punti)


def riquadro(doc, testo: str, sottotitolo: str = "", colore: str | None = None, sfondo: str | None = None,
             livello: int = 2) -> None:
    """Titolo in riquadro con barra laterale (lavorazioni, fasi, schede di macchine e sostanze)."""
    h = doc.add_heading(level=livello)
    _spaziatura(h, 12, 4)
    ombreggia_paragrafo(h, sfondo or TINTA)
    _bordo_paragrafo(h, ("left",), colore or ACCENTO, 36, 6)
    if livello == 2:
        n = _numeri(doc)
        n[1], n[2] = n[1] + 1, 0
        _run(h, f"{n[0]}.{n[1]}  ", colore=TESTO3)
    _run(h, testo, colore=INCHIOSTRO)
    if sottotitolo:
        _run(h, f"   {sottotitolo}", 8.5, TESTO2, carattere=CARATTERE)


def elenco_lettere(doc, voci: list[str], grassetto_iniziale: bool = False) -> None:
    """Elenco a), b), c)..."""
    for i, voce in enumerate(voci):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.7)
        _spaziatura(p, 0, 2)
        lettera = chr(ord("a") + i) if i < 26 else str(i + 1)
        _run(p, f"{lettera})\t", colore=ACCENTO, grassetto=True)
        p.add_run(voce)


def elenco(doc, voci: list[str], numerato: bool = False) -> None:
    for n, voce in enumerate(voci, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        _spaziatura(p, 0, 1.5)
        _run(p, f"{n}.\t" if numerato else "•\t", colore=ACCENTO, grassetto=numerato)
        p.add_run(voce)


def elenco_colonne(doc, voci: list[str], colonne: int = 3) -> None:
    """Elenco puntato breve distribuito su più colonne (per esempio i DPI)."""
    if not voci:
        return
    t = doc.add_table(rows=1, cols=colonne)
    _bordi_tabella(t)
    per_colonna = -(-len(voci) // colonne)
    for n, cella in enumerate(t.rows[0].cells):
        gruppo = voci[n * per_colonna:(n + 1) * per_colonna]
        for i, voce in enumerate(gruppo):
            p = cella.paragraphs[0] if i == 0 else cella.add_paragraph()
            _spaziatura(p, 0, 1)
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            _run(p, "•  ", 8.5, ACCENTO)
            _run(p, voce, 8.5)
    larghezze_colonne(t, [LARGHEZZA_UTILE / colonne] * colonne)
    _spazio_dopo(doc)


def nota(doc, testo: str, avviso: bool = False) -> None:
    """Box evidenziato con barra laterale (note, chiarimenti, avvertenze)."""
    p = doc.add_paragraph()
    _spaziatura(p, 4, 6)
    p.paragraph_format.left_indent = Cm(0.15)
    ombreggia_paragrafo(p, AVVISO_SFONDO if avviso else SUPERFICIE)
    _bordo_paragrafo(p, ("left",), AVVISO if avviso else LINEA_FORTE, 18, 6)
    _run(p, testo, 8.5, AVVISO if avviso else TESTO2)


def testo(doc, contenuto: str, vuoto: str = "Sezione da compilare.") -> None:
    """Testo libero con un paragrafo per riga; se manca, un segnaposto evidenziato."""
    righe = [r.strip() for r in (contenuto or "").strip().split("\n")]
    if not any(righe):
        nota(doc, vuoto, avviso=True)
        return
    for riga in righe:
        if not riga:
            continue
        if riga[:2] in ("- ", "• ", "* "):
            elenco(doc, [riga[2:]])
        else:
            doc.add_paragraph(riga)


def immagine(doc, percorso: str | Path, larghezza_cm: float = 12, didascalia: str = "") -> bool:
    if not percorso or not Path(percorso).is_file():
        return False
    try:
        doc.add_picture(str(percorso), width=Cm(larghezza_cm))
    except Exception:  # noqa: BLE001 - immagine non leggibile: si salta
        return False
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if didascalia:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, didascalia, 8, TESTO3)
    return True


def griglia_immagini(doc, voci: list[tuple[Path, str]], colonne: int = 4, lato_cm: float = 2.6,
                     dimensione_testo: float = 7) -> None:
    """Immagini con didascalia disposte in griglia (cartelli, DPI)."""
    voci = [(p, t) for p, t in voci if Path(p).is_file()]
    if not voci:
        return
    t = doc.add_table(rows=0, cols=colonne)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordi_tabella(t)
    for i, (percorso, didascalia) in enumerate(voci):
        if i % colonne == 0:
            celle = t.add_row().cells
            _non_spezzare(t.rows[-1])
        cella = celle[i % colonne]
        p = cella.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spaziatura(p, 6, 2)
        p.add_run().add_picture(str(percorso), width=Cm(lato_cm))
        d = cella.add_paragraph()
        d.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spaziatura(d, 0, 6)
        _run(d, didascalia[0].upper() + didascalia[1:] if didascalia else "", dimensione_testo, TESTO2)
    _spazio_dopo(doc)


def icone_in_riga(paragrafo, percorsi: list[Path], lato_cm: float = 0.9) -> None:
    for percorso in percorsi:
        if Path(percorso).is_file():
            paragrafo.add_run().add_picture(str(percorso), width=Cm(lato_cm))
            paragrafo.add_run(" ")


def cartello_testo(doc, testo_cartello: str, sfondo: str = "C8102E") -> None:
    """Cartello di solo testo (es. divieto di accesso) come cella colorata."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cella = t.rows[0].cells[0]
    ombreggia_cella(cella, sfondo)
    larghezze_colonne(t, [12])
    p = cella.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spaziatura(p, 10, 10)
    _run(p, testo_cartello.upper(), 16, "FFFFFF", grassetto=True)
    _spazio_dopo(doc, 8)


def tabella_colorata(doc, intestazioni: list[str], righe: list[list[str]], dimensione: float = 8.5,
                     colore_intestazione: str | None = None, larghezze: list[float] | None = None):
    """Tabella leggera: intestazione su fondo tenue ripetuta a ogni pagina, righe a filetti sottili."""
    t = doc.add_table(rows=1, cols=len(intestazioni))
    _bordi_tabella(t, top=(colore_intestazione or ACCENTO, 8), bottom=(LINEA_FORTE, 6), insideH=(LINEA, 4))
    _margini_celle(t)
    _ripeti_intestazione(t.rows[0])
    for cella, titolo in zip(t.rows[0].cells, intestazioni):
        ombreggia_cella(cella, TINTA)
        p = cella.paragraphs[0]
        _spaziatura(p)
        _run(p, titolo, dimensione - 0.5, ACCENTO_SCURO, grassetto=True)
    for riga in righe:
        r = t.add_row()
        _non_spezzare(r)
        for cella, valore in zip(r.cells, riga):
            for n, testo_cella in enumerate(str(valore if valore not in (None, "") else "—").split("\n")):
                testo_cella = testo_cella if testo_cella != " " else ""
                p = cella.paragraphs[0] if n == 0 else cella.add_paragraph()
                _spaziatura(p, 0, 1 if n else 0)
                _run(p, testo_cella, dimensione, TESTO3 if valore in (None, "", "-", "—") else None)
    if not larghezze and intestazioni[0] in ("N.", "Rev."):
        resto = (LARGHEZZA_UTILE - 1.0) / (len(intestazioni) - 1)
        larghezze = [1.0] + [resto] * (len(intestazioni) - 1)
    if larghezze:
        fattore = LARGHEZZA_UTILE / sum(larghezze)
        larghezze_colonne(t, [l * fattore for l in larghezze])
    _spazio_dopo(doc)
    return t


def colora_livello(cella, livello: int) -> None:
    """Cella di un livello di rischio (0 basso - 3 molto alto) nei colori del programma."""
    sfondo, colore = LIVELLI[max(0, min(livello, 3))]
    ombreggia_cella(cella, sfondo)
    for p in cella.paragraphs:
        for r in p.runs:
            r.font.color.rgb = rgb(colore)
    if cella.paragraphs and cella.paragraphs[0].runs:
        cella.paragraphs[0].runs[0].bold = True


def indice(doc) -> None:
    """Indice dei capitoli (campo di Word aggiornato all'apertura)."""
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    _spaziatura(p, 0, 10)
    _bordo_paragrafo(p, ("bottom",), ACCENTO, 12, 4)
    _run(p, "Indice", 17, INCHIOSTRO, carattere=CARATTERE_TITOLI)
    p = doc.add_paragraph()
    campo(p, 'TOC \\o "1-2" \\h \\z \\u', "L'indice si aggiorna all'apertura in Word (o con tasto destro > "
                                          "Aggiorna campo).")
    aggiorna_campi_all_apertura(doc)


def interruzione_pagina(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def firme(doc, voci: list[tuple[str, str]]) -> None:
    """Riquadri delle firme: ruolo, nominativo e spazio per la firma."""
    if not voci:
        return
    t = doc.add_table(rows=0, cols=2)
    _bordi_tabella(t, insideH=(LINEA, 4), bottom=(LINEA, 4), top=(LINEA, 4))
    _margini_celle(t, 0.12, 0.12)
    for ruolo, nome in voci:
        riga = t.add_row()
        _non_spezzare(riga)
        p = riga.cells[0].paragraphs[0]
        _spaziatura(p)
        _run(p, ruolo, 8, TESTO3)
        p = riga.cells[0].add_paragraph()
        _spaziatura(p, 1, 0)
        _run(p, nome or " ", 9.5, INCHIOSTRO, grassetto=True)
        p = riga.cells[1].paragraphs[0]
        _spaziatura(p, 18, 0)
        _bordo_paragrafo(p, ("bottom",), LINEA_FORTE, 4, 1)
        p = riga.cells[1].add_paragraph()
        _spaziatura(p)
        _run(p, "firma", 7, TESTO3)
    larghezze_colonne(t, [8.4, 9.0])
    _spazio_dopo(doc, 8)


def copertina(doc, *, comune: str, titolo: str, riferimento: str, voci: list[tuple[str, str]],
              data_luogo: str, firme: list[tuple[str, str]], impresa: list[str], logo: str = "",
              revisione: str = "") -> None:
    """Copertina: logo, titolo grande con filetto, impresa, dati principali e firme."""
    p = doc.paragraphs[0] if doc.paragraphs and not doc.paragraphs[0].text else doc.add_paragraph()
    _spaziatura(p, 0, 0)
    if logo and Path(logo).is_file():
        try:
            p.add_run().add_picture(logo, height=Cm(1.8))
        except Exception:  # noqa: BLE001 - logo non leggibile
            pass
    p = doc.add_paragraph()
    _spaziatura(p, 70, 6)
    _run(p, riferimento.upper(), 7.5, ACCENTO, grassetto=True)
    p = doc.add_paragraph()
    _spaziatura(p, 0, 0)
    p.paragraph_format.line_spacing = 1.0
    _run(p, titolo[:1].upper() + titolo[1:].lower(), 32, INCHIOSTRO, carattere=CARATTERE_TITOLI)
    p = doc.add_paragraph()
    _spaziatura(p, 2, 20)
    _bordo_paragrafo(p, ("bottom",), ACCENTO, 24, 8)
    righe = [r for r in impresa if r]
    if righe:
        _run(p, righe[0], 15, ACCENTO_SCURO, carattere=CARATTERE_TITOLI)
        for riga in righe[1:]:
            q = doc.add_paragraph()
            _spaziatura(q)
            _run(q, riga, 8.5, TESTO2)
    spazio = doc.add_paragraph()
    _spaziatura(spazio, 0, 18)

    voci = [(e, v) for e, v in voci if v] + ([("Comune", comune)] if comune else [])
    if revisione:
        voci.append(("Revisione", revisione))
    t = doc.add_table(rows=0, cols=2)
    _bordi_tabella(t, insideH=(LINEA, 4), top=(LINEA_FORTE, 6), bottom=(LINEA_FORTE, 6))
    _margini_celle(t, 0.12, 0.1)
    for etichetta, valore in voci:
        celle = t.add_row().cells
        q = celle[0].paragraphs[0]
        _spaziatura(q)
        _run(q, etichetta.upper(), 7.5, TESTO3, grassetto=True)
        q = celle[1].paragraphs[0]
        _spaziatura(q)
        _run(q, valore, 10.5, INCHIOSTRO)
    larghezze_colonne(t, [3.6, LARGHEZZA_UTILE - 3.6])

    p = doc.add_paragraph()
    _spaziatura(p, 26, 6)
    _run(p, data_luogo, 9, TESTO2)
    _firme_copertina(doc, firme)


def _firme_copertina(doc, voci: list[tuple[str, str]]) -> None:
    if not voci:
        return
    t = doc.add_table(rows=1, cols=len(voci))
    _bordi_tabella(t)
    for cella, (ruolo, nome) in zip(t.rows[0].cells, voci):
        p = cella.paragraphs[0]
        _spaziatura(p, 0, 26)
        _run(p, ruolo, 7.5, TESTO3)
        if nome:
            p.add_run("\n")
            _run(p, nome, 9, INCHIOSTRO, grassetto=True)
        q = cella.add_paragraph()
        _spaziatura(q)
        _bordo_paragrafo(q, ("bottom",), LINEA_FORTE, 4, 1)
    larghezze_colonne(t, [LARGHEZZA_UTILE / len(voci)] * len(voci))


def nuova_sezione_verticale(doc) -> None:
    doc.add_section(WD_SECTION.NEW_PAGE)
