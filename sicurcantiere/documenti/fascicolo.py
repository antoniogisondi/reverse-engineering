"""Fascicolo con le caratteristiche dell'opera (art. 91, comma 1, lettera b), e Allegato XVI del D.Lgs. 81/2008).

Stessa impaginazione del PSC: copertina, capitoli I-III dell'Allegato XVI con le loro schede.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document

from sicurcantiere import anagrafiche, analisi
from sicurcantiere.documenti import stile
from sicurcantiere.documenti.comune import data_it, euro, indirizzo_cantiere, salva
from sicurcantiere.documenti.composizione import CONSIGLIATO, OBBLIGATORIO, Manca, segnaposto
from sicurcantiere.repository import Repository

CANTIERE, SOGGETTI, FASCICOLO = "Dati generali", "Soggetti", "Fascicolo"


def genera_fascicolo(repo: Repository, destinazione: str | Path) -> Path:
    with stile.colori_documento(*stile.COLORI_PSC):
        return _genera(repo, destinazione)


def _genera(repo: Repository, destinazione: str | Path) -> Path:
    c = repo.cantiere()
    fasi = repo.elenco("fasi")
    interventi = repo.elenco("interventi")
    elaborati = repo.elenco("elaborati")
    soggetti = [s for s in anagrafiche.soggetti(repo) if s["nome"].strip()]
    periodo = analisi.periodo_lavori(fasi)
    revisioni = repo.elenco("revisioni", documento="FASCICOLO")
    ultima = revisioni[-1] if revisioni else None
    data_documento = data_it(ultima["data"]) if ultima and ultima["data"] else data_it(date.today())
    revisione = f"Rev. {ultima['numero']} del {data_documento}" if ultima else f"Emissione del {data_documento}"
    comune = f"{c['comune']} ({c['provincia']})" if c["provincia"] else c["comune"]
    coordinatore = next((s for s in soggetti if s["ruolo"] in (anagrafiche.CSP, anagrafiche.CSE)), None)
    committente = next((s for s in soggetti if s["ruolo"] == anagrafiche.COMMITTENTE), None)

    doc = Document()
    stile.imposta_stili(doc)
    stile.imposta_pagina(doc, " · ".join(v for v in (c["nome"], comune) if v),
                         (f"Fascicolo dell'opera · {c['nome']}".strip(" ·"), revisione))
    stile.copertina(
        doc,
        comune=comune,
        titolo="Fascicolo dell'opera",
        riferimento="D.Lgs. 9 aprile 2008, n. 81 · art. 91, comma 1, lettera b), e Allegato XVI",
        voci=[("Cantiere", indirizzo_cantiere(c) or Manca("indirizzo del cantiere", CANTIERE)),
              ("Natura dell'opera", c["natura_opera"]),
              ("Committente", committente["nome"] if committente else Manca("committente", SOGGETTI))],
        data_luogo=f"{c['comune'] or ''}, {data_documento}".strip(", "),
        firme=[("Il coordinatore per la sicurezza", coordinatore["nome"] if coordinatore else "")],
        impresa=[c["nome"] or "Opera"],
        revisione=revisione,
    )
    stile.banda(doc, "Revisioni del documento")
    if revisioni:
        stile.tabella_colorata(doc, ["Rev.", "Data", "Descrizione"], [
            [str(r["numero"]), data_it(r["data"]), r["descrizione"]] for r in revisioni], larghezze=[1.2, 2.6, 13.6])
    else:
        doc.add_paragraph(f"Prima emissione del {data_documento}. Il fascicolo è aggiornato in corso d'opera e a "
                          "ogni modifica dell'opera ed è consegnato al committente al termine dei lavori.")

    stile.capitolo(doc, "Descrizione sintetica dell'opera e soggetti coinvolti",
                   "Allegato XVI, capitolo I - scheda I", nuova_pagina=True)
    stile.scheda(doc, [
        ("Opera", c["nome"] or Manca("denominazione dell'opera", CANTIERE)),
        ("Natura dell'opera", c["natura_opera"] or Manca("natura dell'opera", CANTIERE, CONSIGLIATO)),
        ("Indirizzo", indirizzo_cantiere(c) or Manca("indirizzo del cantiere", CANTIERE)),
        ("Importo dei lavori", euro(c["importo_lavori"]) if c["importo_lavori"] else ""),
        ("Inizio dei lavori", data_it(periodo[0]) if periodo else ""),
        ("Fine dei lavori", data_it(periodo[1]) if periodo else ""),
    ])
    stile.banda(doc, "Descrizione sintetica dell'opera")
    stile.testo(doc, c["descrizione"], "descrizione sintetica dell'opera", CANTIERE, OBBLIGATORIO)
    stile.banda(doc, "Soggetti")
    if soggetti:
        stile.tabella_colorata(doc, ["Ruolo", "Nominativo", "Recapiti"], [
            [s["ruolo"], s["nome"], " · ".join(v for v in (s["indirizzo"], s["telefono"], s["email"]) if v)]
            for s in soggetti], larghezze=[4.6, 4.6, 8.2])
    else:
        stile.testo(doc, "", "committente e coordinatori", SOGGETTI, OBBLIGATORIO)

    stile.capitolo(doc, "Misure preventive e protettive in dotazione dell'opera e ausiliarie",
                   "Allegato XVI, capitolo II")
    stile.sezione(doc, "Interventi di manutenzione e relative misure", "scheda II-1")
    if interventi:
        for parte in dict.fromkeys(i["parte_opera"] for i in interventi):
            stile.riquadro(doc, parte or "Opera", "parte d'opera", stile.TESTO3, stile.SUPERFICIE, livello=3)
            stile.tabella_colorata(doc, ["Intervento", "Periodicità", "Rischi individuati",
                                         "Misure in dotazione dell'opera", "Misure ausiliarie"], [
                [i["intervento"], i["periodicita"], i["rischi"], i["misure_in_dotazione"], i["misure_ausiliarie"]]
                for i in interventi if i["parte_opera"] == parte], 8, larghezze=[3.6, 2.0, 3.8, 4.0, 4.0])
    else:
        stile.testo(doc, "", "interventi di manutenzione dell'opera, con rischi e misure", FASCICOLO,
                    OBBLIGATORIO)

    stile.sezione(doc, "Misure preventive e protettive in dotazione dell'opera", "scheda II-3")
    dotazioni = sorted({i["misure_in_dotazione"].strip() for i in interventi if i["misure_in_dotazione"].strip()})
    if dotazioni:
        stile.tabella_colorata(doc, ["Misura in dotazione", "Verifiche e controlli da effettuare", "Periodicità"],
                               [[m, " ", " "] for m in dotazioni], larghezze=[6.4, 7.4, 3.6])
        stile.nota(doc, "Le verifiche, i controlli e le relative periodicità sono completati sulla base delle "
                        "indicazioni dei fabbricanti e degli installatori delle misure in dotazione.")
    else:
        stile.testo(doc, "", "misure preventive e protettive in dotazione dell'opera (ad esempio ancoraggi e "
                             "linee vita)", FASCICOLO, CONSIGLIATO)

    stile.capitolo(doc, "Riferimenti alla documentazione di supporto", "Allegato XVI, capitolo III")
    if elaborati:
        stile.tabella_colorata(doc, ["Elaborato", "Autore", "Data", "Collocazione"], [
            [e["titolo"], e["autore"], data_it(e["data"]) if e["data"] else "", e["collocazione"]]
            for e in elaborati], larghezze=[6.4, 4.0, 2.4, 4.6])
    else:
        stile.testo(doc, "", "elaborati tecnici dell'opera (progetto architettonico, strutturale, impianti) e loro "
                             "collocazione", FASCICOLO, CONSIGLIATO)

    stile.capitolo(doc, "Firme")
    p = doc.add_paragraph(f"{c['comune'] or ''}, {data_documento}".strip(", "))
    stile._spaziatura(p, 0, 8)
    firme = [("Il coordinatore per la sicurezza", coordinatore["nome"] if coordinatore else "")]
    if committente:
        firme.append(("Per ricevuta, il committente", committente["nome"]))
    stile.firme(doc, firme)
    if coordinatore is None:
        segnaposto(doc, doc.add_paragraph(), Manca("coordinatore per la sicurezza (CSP o CSE)", SOGGETTI,
                                                   OBBLIGATORIO))
    return salva(doc, destinazione)
