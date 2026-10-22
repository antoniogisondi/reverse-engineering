"""Verbale di verifica dell'idoneità tecnico-professionale di un'impresa o di un lavoratore autonomo
(art. 90, comma 9, art. 97, comma 2, Allegato XVII), con patente a crediti (art. 27) e DURC di congruità."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document

from sicurcantiere import anagrafiche
from sicurcantiere.documenti import stile
from sicurcantiere.documenti.comune import data_it, indirizzo_cantiere, salva
from sicurcantiere.documenti.composizione import OBBLIGATORIO, Manca
from sicurcantiere.repository import Repository

IMPRESE = "Impresa"


def genera_verifica_idoneita(repo: Repository, impresa_id: int, destinazione: str | Path) -> Path:
    impresa = repo.leggi("imprese", impresa_id)
    if impresa is None:
        raise ValueError("Impresa inesistente")
    c = repo.cantiere()
    soggetti = anagrafiche.soggetti(repo)
    committente = next((s for s in soggetti if s["ruolo"] == anagrafiche.COMMITTENTE and s["nome"].strip()), None)
    responsabile = next((s for s in soggetti if s["ruolo"] == anagrafiche.RESPONSABILE_LAVORI and s["nome"].strip()),
                        None)
    affidataria = repo.leggi("imprese", impresa["affidataria_id"]) if impresa.get("affidataria_id") else None
    autonomo = anagrafiche.autonomo(impresa)
    comune = f"{c['comune']} ({c['provincia']})" if c["provincia"] else c["comune"]
    data_verifica = data_it(impresa["idoneita_verificata"]) if impresa.get("idoneita_verificata") else data_it(date.today())

    doc = Document()
    stile.imposta_stili(doc)
    stile.imposta_pagina(doc, " · ".join(v for v in (c["nome"], comune) if v),
                         ("Verifica dell'idoneità tecnico-professionale", data_verifica))
    for sezione in doc.sections:
        sezione.different_first_page_header_footer = False

    p = doc.add_paragraph()
    stile._run(p, "D.LGS. 9 APRILE 2008, N. 81 · ART. 90, COMMA 9, ART. 97, COMMA 2, E ALLEGATO XVII", 7.5,
               stile.ACCENTO, grassetto=True)
    p = doc.add_paragraph()
    stile._bordo_paragrafo(p, ("bottom",), stile.ACCENTO, 18, 6)
    stile._run(p, "Verifica dell'idoneità tecnico-professionale", 20, stile.INCHIOSTRO, carattere=stile.CARATTERE_TITOLI)
    p = doc.add_paragraph()
    stile._spaziatura(p, 4, 10)
    stile._run(p, ("del lavoratore autonomo " if autonomo else "dell'impresa ") + impresa["ragione_sociale"], 12,
               stile.ACCENTO_SCURO, carattere=stile.CARATTERE_TITOLI)

    stile.banda(doc, "Cantiere")
    stile.scheda(doc, [("Opera", c["nome"]), ("Indirizzo", indirizzo_cantiere(c)),
                       ("Committente", committente["nome"] if committente else ""),
                       ("Responsabile dei lavori", responsabile["nome"] if responsabile else "")])
    stile.banda(doc, "Lavoratore autonomo" if autonomo else "Impresa")
    stile.scheda(doc, [
        ("Ragione sociale", impresa["ragione_sociale"]), ("Ruolo nel cantiere", impresa["tipo"]),
        ("Datore di lavoro", impresa["datore_lavoro"]), ("Sede legale", anagrafiche.sede(impresa)),
        ("Codice fiscale", impresa["codice_fiscale"]), ("Partita IVA", impresa["partita_iva"]),
        ("Lavori affidati", impresa["tipologia_lavori"]),
        ("Impresa affidataria", affidataria["ragione_sociale"] if affidataria else ""),
    ])

    stile.banda(doc, "Patente a crediti e regolarità")
    stile.scheda(doc, [
        ("Patente a crediti (art. 27)", anagrafiche.descrizione_patente(impresa) or
         Manca(f"patente a crediti di {impresa['ragione_sociale']}", IMPRESE, OBBLIGATORIO)),
        ("Verificata sul portale INL il", data_it(impresa["patente_verificata"]) if impresa.get("patente_verificata")
         else ""),
        ("DURC valido fino al", data_it(impresa["durc_scadenza"]) if impresa.get("durc_scadenza") else ""),
        ("DURC di congruità (D.M. 143/2021)", " · ".join(v for v in (
            impresa.get("durc_congruita"), data_it(impresa["durc_congruita_data"])
            if impresa.get("durc_congruita_data") else "") if v)),
    ], 6.0)
    if anagrafiche.patente_regolare(impresa) is False:
        stile.nota(doc, f"Patente a crediti assente o con meno di {anagrafiche.CREDITI_MINIMI} crediti: l'impresa non "
                        "può operare nel cantiere (art. 27).", avviso=True)

    stile.banda(doc, "Documenti verificati (Allegato XVII, punto " + ("2)" if autonomo else "1)"))
    voci = anagrafiche.leggi_idoneita(impresa)
    righe = [[str(n), testo, voci[k]["esito"] or "Da verificare", voci[k]["nota"]]
             for n, (k, testo) in enumerate(anagrafiche.voci_idoneita(impresa), 1)]
    t = stile.tabella_colorata(doc, ["N.", "Documento", "Esito", "Note"], righe, 8, larghezze=[0.8, 9.0, 2.4, 5.2])
    for riga, (k, _) in zip(t.rows[1:], anagrafiche.voci_idoneita(impresa)):
        esito = voci[k]["esito"]
        stile.colora_livello(riga.cells[2], 0 if esito in ("Acquisito", "Non pertinente") else 3 if esito == "Mancante"
                             else 1)

    esito = anagrafiche.esito_idoneita(impresa)
    stile.banda(doc, "Esito")
    p = doc.add_paragraph()
    stile._run(p, {"Idonea": "Idoneità tecnico-professionale verificata: la documentazione richiesta è completa.",
                   "Documentazione incompleta": "Documentazione incompleta: l'impresa non può iniziare i lavori finché "
                                                "non sono acquisiti i documenti mancanti.",
                   }.get(esito, "Verifica non ancora eseguita."), 10.5,
               stile.LIVELLI[0][1] if esito == "Idonea" else stile.AVVISO, grassetto=True)

    p = doc.add_paragraph(f"{c['comune'] or ''}, {data_verifica}".strip(", "))
    stile._spaziatura(p, 14, 8)
    if anagrafiche.autonomo(impresa) or not affidataria:
        verificatore = ("Il committente o il responsabile dei lavori",
                        (responsabile or committente or {}).get("nome", ""))
    else:
        verificatore = (f"Il datore di lavoro dell'impresa affidataria {affidataria['ragione_sociale']}",
                        affidataria["datore_lavoro"])
    stile.firme(doc, [verificatore, ("Per l'impresa verificata, il datore di lavoro" if not autonomo else
                                     "Il lavoratore autonomo", impresa["datore_lavoro"] or impresa["ragione_sociale"])])
    return salva(doc, destinazione)
