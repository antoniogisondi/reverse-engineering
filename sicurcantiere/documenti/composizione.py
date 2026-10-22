"""Composizione dei documenti: dati mancanti, capitoli, bozza e versione definitiva.

Ogni dato che manca nel documento viene segnato con `Manca` (cosa manca, in quale pagina del programma si
compila, quanto è importante). Nella bozza compare evidenziato in giallo, con un commento di Word che dice
dove compilarlo, e in un riepilogo «Dati da completare» prima dell'indice; nella versione definitiva diventa
una riga da compilare a penna. I dati facoltativi non si stampano: compaiono solo nel riepilogo.

Durante una composizione (`in_composizione`) il documento viene diviso in capitoli: ognuno si può
togliere, sostituire con un testo proprio o completare in fondo (le personalizzazioni restano quando si
ricompone), e il testo di ogni capitolo si conserva per confrontare due revisioni.
"""

from __future__ import annotations

import hashlib
import io
import re
from contextlib import contextmanager
from contextvars import ContextVar
from dataclasses import dataclass, field
from typing import Any, Iterator

from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

OBBLIGATORIO, CONSIGLIATO, FACOLTATIVO = "obbligatorio", "consigliato", "facoltativo"
LIVELLI = {
    OBBLIGATORIO: "Obbligatorio per legge",
    CONSIGLIATO: "Consigliato",
    FACOLTATIVO: "Facoltativo",
}

# Modi di personalizzare un capitolo.
AGGIUNGI, SOSTITUISCI = "aggiungi", "sostituisci"

# Stili che riconoscono, anche dopo le modifiche in Word, ciò che appartiene solo alla bozza.
STILE_DATO = "SCDaCompilare"          # dato mancante in una riga (carattere)
STILE_TESTO = "SCDaCompilareTesto"    # sezione di testo mancante (carattere)
STILE_BOZZA = "SCSoloBozza"           # paragrafi della sola bozza: riepilogo e filigrana
AUTORE = "SicurCantiere"
RIGA_PENNA = "_" * 28
RIGHE_PENNA = 3
LUNGHEZZA_RIGA = 92


@dataclass(frozen=True)
class Manca:
    """Dato mancante: cosa manca, sezione del programma in cui si compila, importanza."""
    dato: str
    sezione: str = ""
    livello: str = OBBLIGATORIO


@dataclass
class Contesto:
    """Stato della generazione di un documento."""
    composizione: bool = False  # composizione salvata: capitoli, riepilogo e filigrana
    personalizzazioni: dict[str, dict[str, Any]] = field(default_factory=dict)
    mancanti: list[dict[str, Any]] = field(default_factory=list)
    capitoli: list[dict[str, Any]] = field(default_factory=list)
    confini: list[Any] = field(default_factory=list)  # elementi che chiudono un capitolo (parti, indice)
    ancora: Any = None  # elemento prima del quale va il riepilogo dei dati mancanti
    capitolo: str = ""  # capitolo in corso, per il riepilogo
    _elementi_mancanti: list[Any] = field(default_factory=list)


_corrente: ContextVar[Contesto | None] = ContextVar("composizione", default=None)


@contextmanager
def in_composizione(personalizzazioni: dict[str, dict[str, Any]] | None = None) -> Iterator[Contesto]:
    """I documenti generati qui dentro diventano una composizione: si legge il contesto alla fine."""
    contesto_ = Contesto(composizione=True, personalizzazioni=dict(personalizzazioni or {}))
    gettone = _corrente.set(contesto_)
    try:
        yield contesto_
    finally:
        _corrente.reset(gettone)


def contesto(doc) -> Contesto:
    c = getattr(doc, "_sc_contesto", None)
    if c is None:
        c = _corrente.get() or Contesto()
        doc._sc_contesto = c
    return c


# --- dati mancanti ---

def _stile_carattere(doc, nome: str):
    stili = doc.styles
    if nome in [s.name for s in stili]:
        return stili[nome]
    s = stili.add_style(nome, WD_STYLE_TYPE.CHARACTER)
    s.hidden = False
    s.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return s


def _stile_paragrafo(doc, nome: str):
    stili = doc.styles
    if nome in [s.name for s in stili]:
        return stili[nome]
    s = stili.add_style(nome, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = stili["Normal"]
    return s


def _commento(doc, run, m: Manca) -> None:
    testo = LIVELLI.get(m.livello, m.livello) + "."
    if m.sezione:
        testo += f" Si compila in SicurCantiere, pagina o scheda «{m.sezione}»."
    try:
        doc.add_comment(run, text=testo, author=AUTORE, initials="SC")
    except Exception:  # noqa: BLE001 - un commento non riuscito non deve fermare il documento
        pass


def segna(doc, m: Manca, elemento: Any = None) -> None:
    c = contesto(doc)
    c.mancanti.append({"dato": m.dato, "sezione": m.sezione, "livello": m.livello, "capitolo": c.capitolo})
    c._elementi_mancanti.append(elemento)


def segnaposto(doc, paragrafo, m: Manca):
    """Dato mancante dentro una riga: «[Da compilare: ...]» evidenziato, con commento. I dati facoltativi
    si annotano soltanto e non compaiono nel documento."""
    if m.livello == FACOLTATIVO:
        segna(doc, m)
        return None
    run = paragrafo.add_run(f"[Da compilare: {m.dato}]")
    run.style = _stile_carattere(doc, STILE_DATO)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    segna(doc, m, run._r)
    _commento(doc, run, m)
    return run


def blocco(doc, m: Manca, guida: str = ""):
    """Sezione di testo mancante: un paragrafo evidenziato con la guida alla compilazione."""
    if m.livello == FACOLTATIVO:
        segna(doc, m)
        return None
    p = doc.add_paragraph()
    run = p.add_run(f"[Da compilare: {guida or m.dato}]")
    run.style = _stile_carattere(doc, STILE_TESTO)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    segna(doc, m, run._r)
    _commento(doc, run, m)
    return p


def mancanti_presenti(doc) -> list[dict[str, Any]]:
    """Dati mancanti rimasti nel documento (non quelli dei capitoli tolti o sostituiti)."""
    c = contesto(doc)
    radice = doc.element
    esclusi = {cap["titolo"] for cap in c.capitoli if cap.get("escluso") or cap.get("modo") == SOSTITUISCI}
    restano = []
    for voce, elemento in zip(c.mancanti, c._elementi_mancanti):
        if elemento is not None:
            while elemento.getparent() is not None:
                elemento = elemento.getparent()
            if elemento is not radice:  # tolto con il suo capitolo
                continue
        elif voce["capitolo"] and _senza_numero(voce["capitolo"]) in esclusi:
            continue
        restano.append(voce)
    return restano


# --- capitoli ---

def _senza_numero(titolo: str) -> str:
    return re.sub(r"^[\d.\s]+", "", titolo).strip()


def chiave_capitolo(titolo: str) -> str:
    testo = _senza_numero(titolo).lower()
    return re.sub(r"[^a-z0-9àèéìòù]+", "-", testo).strip("-") or "capitolo"


def _e_titolo(elemento) -> bool:
    if elemento.tag != qn("w:p"):
        return False
    stile = elemento.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    return stile is not None and stile.get(qn("w:val")) in ("Heading1", "Titolo1")


def _testo_elemento(elemento) -> str:
    if elemento.tag == qn("w:tbl"):
        righe = []
        for tr in elemento.iter(qn("w:tr")):
            celle = ["".join(t.text or "" for t in tc.iter(qn("w:t"))).strip() for tc in tr.iter(qn("w:tc"))]
            righe.append(" | ".join(celle))
        return "\n".join(righe)
    return "".join(t.text or "" for t in elemento.iter(qn("w:t")))


def impronta(testo: str) -> str:
    return hashlib.sha1(testo.encode("utf-8")).hexdigest()[:12]


def _dividi_capitoli(doc, c: Contesto) -> list[dict[str, Any]]:
    corpo = list(doc.element.body)
    confini = {id(e) for e in c.confini}
    capitoli, corrente = [], None
    for e in corpo:
        if e.tag == qn("w:sectPr"):
            break
        if _e_titolo(e):
            corrente = {"titolo": _testo_elemento(e).strip(), "titolo_el": e, "corpo": []}
            capitoli.append(corrente)
        elif id(e) in confini or (e.tag == qn("w:p") and _stile_paragrafo_di(e) == STILE_BOZZA):
            corrente = None
        elif corrente is not None:
            corrente["corpo"].append(e)
    visti: dict[str, int] = {}
    for cap in capitoli:
        chiave = chiave_capitolo(cap["titolo"])
        visti[chiave] = visti.get(chiave, 0) + 1
        cap["chiave"] = chiave if visti[chiave] == 1 else f"{chiave}-{visti[chiave]}"
    return capitoli


def _stile_paragrafo_di(elemento) -> str:
    stile = elemento.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    return stile.get(qn("w:val")) if stile is not None else ""


def _riferimento(elementi: list) -> list:
    """Il riferimento di legge sotto il titolo (paragrafo piccolo che segue il titolo) resta al suo posto."""
    if elementi and elementi[0].tag == qn("w:p"):
        dimensioni = elementi[0].findall(f".//{qn('w:sz')}")
        if dimensioni and all(d.get(qn("w:val")) == "16" for d in dimensioni):
            return elementi[:1]
    return []


def _aggiungi_testo(doc, prima_di, testo: str) -> None:
    """Testo scritto dal consulente, inserito prima dell'elemento indicato (o in fondo)."""
    from sicurcantiere.documenti import stile

    corpo = doc.element.body
    presenti = {id(e) for e in corpo}
    stile.testo(doc, testo, "")
    nuovi = [e for e in corpo if id(e) not in presenti and e.tag != qn("w:sectPr")]
    if prima_di is not None:
        for e in nuovi:
            prima_di.addprevious(e)


def _applica_personalizzazioni(doc, c: Contesto) -> None:
    capitoli = _dividi_capitoli(doc, c)
    for cap in capitoli:
        testo = "\n".join(t for t in (_testo_elemento(e).strip() for e in cap["corpo"]) if t)
        # «generato»: composto dal programma (per l'editor dei capitoli); «testo»: quello stampato (per il
        # confronto tra revisioni)
        voce = {"chiave": cap["chiave"], "titolo": _senza_numero(cap["titolo"]), "generato": testo, "testo": testo,
                "impronta": impronta(testo), "escluso": False, "modo": "", "cambiato": False}
        p = c.personalizzazioni.get(cap["chiave"])
        if p:
            voce["cambiato"] = bool(p.get("impronta")) and p["impronta"] != voce["impronta"]
            if p.get("escluso"):
                voce["escluso"] = True
                for e in [cap["titolo_el"], *cap["corpo"]]:
                    e.getparent().remove(e)
            elif (p.get("testo") or "").strip():
                voce["modo"] = p.get("modo") or AGGIUNGI
                ultimo = cap["corpo"][-1] if cap["corpo"] else cap["titolo_el"]
                successivo = ultimo.getnext()
                proprio = "\n".join(r.strip()[2:] if r.strip()[:2] in ("- ", "• ", "* ") else r.strip()
                                     for r in p["testo"].strip().splitlines() if r.strip())
                if voce["modo"] == SOSTITUISCI:
                    restano = _riferimento(cap["corpo"])
                    for e in cap["corpo"][len(restano):]:
                        e.getparent().remove(e)
                    voce["testo"] = "\n".join(t for t in [*(_testo_elemento(e).strip() for e in restano), proprio]
                                               if t)
                else:
                    voce["testo"] = "\n".join(t for t in (testo, proprio) if t)
                if successivo is not None and successivo.tag == qn("w:sectPr"):
                    successivo = None
                _aggiungi_testo(doc, successivo, p["testo"])
        c.capitoli.append(voce)


# --- solo bozza: riepilogo e filigrana ---

def _paragrafo_bozza(doc, prima_di, nuova_pagina: bool = False):
    p = doc.add_paragraph(style=_stile_paragrafo(doc, STILE_BOZZA))
    if nuova_pagina:
        p.paragraph_format.page_break_before = True
    if prima_di is not None:
        prima_di.addprevious(p._p)
    return p


def _riepilogo(doc, c: Contesto, mancanti: list[dict[str, Any]]) -> None:
    """Pagina «Dati da completare» con l'aspetto dell'indice: titolo con filetto, gruppi come le bande dei
    capitoli, voci come gli elenchi puntati."""
    from docx.shared import Cm

    from sicurcantiere.documenti import stile

    ancora = c.ancora
    if ancora is None:
        ancora = next((e for e in doc.element.body if _e_titolo(e)), None)
    p = _paragrafo_bozza(doc, ancora, nuova_pagina=True)
    stile._spaziatura(p, 0, 10)
    stile._bordo_paragrafo(p, ("bottom",), stile.ACCENTO, 12, 4)
    stile._run(p, "Dati da completare", 17, stile.INCHIOSTRO, carattere=stile.CARATTERE_TITOLI)
    p = _paragrafo_bozza(doc, ancora)
    stile.ombreggia_paragrafo(p, stile.AVVISO_SFONDO)
    stile._bordo_paragrafo(p, ("left",), stile.AVVISO, 18, 6)
    stile._spaziatura(p, 4, 6)
    stile._run(p, "Pagina presente solo nella bozza: nella versione emessa non compare. I dati evidenziati in giallo "
                  "si compilano nelle pagine del programma indicate qui sotto; quelli rimasti vuoti all'emissione "
                  "diventano righe da compilare a mano.", 8.5, stile.AVVISO)
    if not mancanti:
        stile._run(_paragrafo_bozza(doc, ancora), "Nessun dato mancante: il documento è completo.", colore=stile.TESTO2)
        return
    for livello in (OBBLIGATORIO, CONSIGLIATO, FACOLTATIVO):
        voci = [m for m in mancanti if m["livello"] == livello]
        if not voci:
            continue
        p = _paragrafo_bozza(doc, ancora)  # come stile.banda
        stile._spaziatura(p, 10, 4)
        p.paragraph_format.keep_with_next = True
        stile._bordo_paragrafo(p, ("bottom",), stile.LINEA, 4, 2)
        r = stile._run(p, f"{LIVELLI[livello]}: {len(voci)}", 8, stile.AVVISO if livello == OBBLIGATORIO
                       else stile.ACCENTO_SCURO, grassetto=True)
        r.font.all_caps = True
        for m in voci:
            dove = " · ".join(v for v in (m["capitolo"] and f"cap. {m['capitolo']}",
                                          m["sezione"] and f"si compila in «{m['sezione']}»") if v)
            p = _paragrafo_bozza(doc, ancora)  # come stile.elenco
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            stile._spaziatura(p, 0, 1.5)
            stile._run(p, "•\t", colore=stile.ACCENTO, grassetto=False)
            stile._run(p, m["dato"][:1].upper() + m["dato"][1:])
            if dove:
                stile._run(p, f"  ·  {dove}", 8, stile.TESTO3)


FILIGRANA = (
    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">'
    '<w:pict>'
    '<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" '
    'path="m@7,l@8,m@5,21600l@6,21600e">'
    '<v:formulas><v:f eqn="sum #0 0 10800"/><v:f eqn="prod #0 2 1"/><v:f eqn="sum 21600 0 @1"/>'
    '<v:f eqn="sum 0 0 @2"/><v:f eqn="sum 21600 0 @3"/><v:f eqn="if @0 @3 0"/><v:f eqn="if @0 21600 @1"/>'
    '<v:f eqn="if @0 0 @2"/><v:f eqn="if @0 @4 21600"/><v:f eqn="mid @5 @6"/><v:f eqn="mid @8 @5"/>'
    '<v:f eqn="mid @7 @8"/><v:f eqn="mid @6 @7"/><v:f eqn="sum @6 0 @5"/></v:formulas>'
    '<v:path textpathok="t" o:connecttype="custom" '
    'o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" o:connectangles="270,180,90,0"/>'
    '<v:textpath on="t" fitshape="t"/><o:lock v:ext="edit" text="t" shapetype="t"/></v:shapetype>'
    '<v:shape id="SicurCantiereBozza{n}" o:spid="_x0000_s20{n:02d}" type="#_x0000_t136" '
    'style="position:absolute;margin-left:0;margin-top:0;width:430pt;height:140pt;rotation:315;'
    'z-index:-251655168;mso-position-horizontal:center;mso-position-horizontal-relative:margin;'
    'mso-position-vertical:center;mso-position-vertical-relative:margin" o:allowincell="f" '
    'fillcolor="#d0d6dc" stroked="f">'
    '<v:fill opacity=".45"/><v:textpath style="font-family:&quot;Ubuntu&quot;;font-size:1pt" string="BOZZA"/>'
    '</v:shape></w:pict></w:r>'
)


def _filigrana(doc) -> None:
    n = 0
    for i, sezione in enumerate(doc.sections):
        intestazioni = [sezione.header]
        if sezione.different_first_page_header_footer:
            intestazioni.append(sezione.first_page_header)
        for intestazione in intestazioni:
            if i and intestazione.is_linked_to_previous:
                continue  # la sezione usa l'intestazione della precedente, che ha già la filigrana
            n += 1
            p = intestazione.add_paragraph(style=_stile_paragrafo(doc, STILE_BOZZA))
            p.paragraph_format.space_after = 0
            p._p.append(parse_xml(FILIGRANA.replace("{n}", str(n)).replace("{n:02d}", f"{n:02d}")))


def finalizza(doc) -> None:
    """Chiusura della composizione (prima del salvataggio): capitoli personalizzati, riepilogo dei dati
    mancanti e filigrana BOZZA."""
    c = contesto(doc)
    if not c.composizione:
        return
    _applica_personalizzazioni(doc, c)
    mancanti = mancanti_presenti(doc)
    c.mancanti, c._elementi_mancanti = mancanti, [None] * len(mancanti)
    _riepilogo(doc, c, mancanti)
    _filigrana(doc)


# --- dalla bozza alla versione definitiva ---

def _ids_stili(documento) -> dict[str, str]:
    return {s.name: s.style_id for s in documento.styles}


def definitivo(contenuto: bytes) -> bytes:
    """Versione da emettere: via filigrana, riepilogo e commenti di SicurCantiere; i dati mancanti
    diventano righe da compilare a mano. Funziona anche sul file modificato in Word."""
    from docx import Document

    documento = Document(io.BytesIO(contenuto))
    ids = _ids_stili(documento)
    stile_bozza = ids.get(STILE_BOZZA, STILE_BOZZA)
    stile_dato = ids.get(STILE_DATO, STILE_DATO)
    stile_testo = ids.get(STILE_TESTO, STILE_TESTO)

    parti = [documento.element.body]
    for sezione in documento.sections:
        for parte in (sezione.header, sezione.first_page_header, sezione.even_page_header):
            if not parte.is_linked_to_previous:
                parti.append(parte._element)

    for radice in parti:
        for p in list(radice.iter(qn("w:p"))):
            if _stile_paragrafo_di(p) == stile_bozza and p.getparent() is not None:
                p.getparent().remove(p)
        for r in list(radice.iter(qn("w:r"))):
            rpr = r.find(qn("w:rPr"))
            stile_run = rpr.find(qn("w:rStyle")) if rpr is not None else None
            valore = stile_run.get(qn("w:val")) if stile_run is not None else ""
            if valore not in (stile_dato, stile_testo):
                continue
            rpr.remove(stile_run)
            for evidenza in rpr.findall(qn("w:highlight")):
                rpr.remove(evidenza)
            for t in r.findall(qn("w:t")):
                r.remove(t)
            righe = 1 if valore == stile_dato else RIGHE_PENNA
            lunghezza = len(RIGA_PENNA) if valore == stile_dato else LUNGHEZZA_RIGA
            for i in range(righe):
                if i:
                    r.append(OxmlElement("w:br"))
                t = OxmlElement("w:t")
                t.text = "_" * lunghezza
                r.append(t)

    _togli_commenti(documento)
    uscita = io.BytesIO()
    documento.save(uscita)
    return uscita.getvalue()


def _togli_commenti(documento) -> None:
    parte = getattr(documento.part, "_comments_part", None)
    try:
        commenti = documento.part._comments_part.element if parte is not None else None
    except Exception:  # noqa: BLE001
        commenti = None
    if commenti is None:
        return
    nostri = {c.get(qn("w:id")) for c in commenti.findall(qn("w:comment")) if c.get(qn("w:author")) == AUTORE}
    if not nostri:
        return
    corpo = documento.element.body
    for nome in ("w:commentRangeStart", "w:commentRangeEnd"):
        for e in list(corpo.iter(qn(nome))):
            if e.get(qn("w:id")) in nostri:
                e.getparent().remove(e)
    for riferimento in list(corpo.iter(qn("w:commentReference"))):
        if riferimento.get(qn("w:id")) in nostri:
            run = riferimento.getparent()
            run.getparent().remove(run)
    for c in list(commenti.findall(qn("w:comment"))):
        if c.get(qn("w:id")) in nostri:
            commenti.remove(c)


def leggi_mancanti(contenuto: bytes) -> int:
    """Quanti dati da compilare restano evidenziati in un file (anche dopo le modifiche in Word)."""
    from docx import Document

    documento = Document(io.BytesIO(contenuto))
    ids = _ids_stili(documento)
    cercati = {ids.get(STILE_DATO, STILE_DATO), ids.get(STILE_TESTO, STILE_TESTO)}
    return sum(1 for s in documento.element.body.iter(qn("w:rStyle")) if s.get(qn("w:val")) in cercati)
