"""Funzioni comuni per la generazione dei documenti DOCX."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from sicurcantiere import analisi
from sicurcantiere.repository import Repository


def euro(valore: float) -> str:
    testo = f"{valore:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"€ {testo}"


def data_it(valore: str | date | None) -> str:
    d = analisi.data(valore) if isinstance(valore, str) or valore is None else valore
    return d.strftime("%d/%m/%Y") if d else "-"


def nome_fase(f: dict[str, Any]) -> str:
    return f"{f['codice']} {f['descrizione']}".strip()


def indirizzo_cantiere(c: dict[str, Any]) -> str:
    comune = f"{c['comune']} ({c['provincia']})" if c["provincia"] else c["comune"]
    return ", ".join(v for v in (c["indirizzo"], comune) if v)


def ombreggia(cella, colore: str) -> None:
    proprieta = cella._tc.get_or_add_tcPr()
    sfondo = OxmlElement("w:shd")
    sfondo.set(qn("w:val"), "clear")
    sfondo.set(qn("w:color"), "auto")
    sfondo.set(qn("w:fill"), colore)
    proprieta.append(sfondo)


def tabella(doc, intestazioni: list[str], righe: list[list[str]], dimensione: float | None = None):
    """Tabella con l'impostazione grafica dei documenti (intestazione ripetuta, filetti sottili)."""
    from sicurcantiere.documenti import stile

    return stile.tabella_colorata(doc, intestazioni, [[str(v) for v in r] for r in righe], dimensione or 8.5)


def tabella_voci(doc, voci: list[tuple[str, str]]) -> None:
    """Dati etichetta-valore; le voci vuote non si stampano."""
    from sicurcantiere.documenti import stile

    stile.scheda(doc, voci)


def testo_o_segnaposto(doc, testo: str) -> None:
    if testo.strip():
        for paragrafo in testo.strip().split("\n"):
            doc.add_paragraph(paragrafo)
    else:
        doc.add_paragraph().add_run("[Sezione da compilare]").italic = True


def nuovo_documento():
    doc = Document()
    for sezione in doc.sections:
        sezione.left_margin = sezione.right_margin = Cm(2)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    return doc


def frontespizio(doc, repo: Repository, titolo: str, riferimento: str, sigla: str,
                 voci_extra: list[tuple[str, str]] = (), logo: str = "") -> None:
    """Frontespizio con dati del cantiere e tabella delle revisioni del documento."""
    c = repo.cantiere()
    if logo and Path(logo).is_file():
        try:
            doc.add_paragraph().add_run().add_picture(logo, height=Cm(2.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:  # noqa: BLE001 - immagine non leggibile: il frontespizio resta senza logo
            pass
    t = doc.add_heading(titolo, 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(riferimento)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    for etichetta, valore in [
        ("Opera", c["nome"]),
        ("Natura dell'opera", c["natura_opera"]),
        ("Indirizzo del cantiere", indirizzo_cantiere(c)),
        *voci_extra,
        ("Data", data_it(date.today())),
    ]:
        p = doc.add_paragraph()
        p.add_run(f"{etichetta}: ").bold = True
        p.add_run(valore or "-")

    revisioni = repo.elenco("revisioni", documento=sigla)
    if revisioni:
        doc.add_paragraph()
        doc.add_paragraph().add_run("Revisioni del documento").bold = True
        tabella(doc, ["Rev.", "Data", "Descrizione"], [
            [str(r["numero"]), data_it(r["data"]), r["descrizione"]] for r in revisioni
        ])
    doc.add_page_break()


def firma(doc, ruolo: str) -> None:
    doc.add_paragraph()
    doc.add_paragraph(ruolo)
    doc.add_paragraph("_______________________________")
