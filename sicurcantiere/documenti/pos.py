"""Generazione del Piano Operativo di Sicurezza (Allegato XV, punto 3.2.1)."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

from sicurcantiere import anagrafiche, analisi, segnaletica
from sicurcantiere.documenti import stile
from sicurcantiere.documenti.comune import (
    data_it,
    euro,
    firma,
    indirizzo_cantiere,
    nome_fase,
    testo_o_segnaposto,
)
from sicurcantiere.documenti.sostanze import aggiungi_sostanze, sostanze_impresa
from sicurcantiere.documenti.valutazione import aggiungi_valutazione, lavorazioni_valutazione, valutazioni_impresa
from sicurcantiere.repository import Repository
from sicurcantiere.struttura import MACCHINA, struttura_fase
from sicurcantiere.valutazioni import METODI
from sicurcantiere.valutazioni.automatiche import da_eseguire

# Sezioni a testo libero del POS, salvate per impresa.
SEZIONI_POS: list[tuple[str, str]] = [
    ("attivita", "Attività svolte in cantiere, modalità organizzative e turni di lavoro"),
    ("sostanze", "Sostanze pericolose: note integrative all'elenco"),
    ("rumore", "Esito della valutazione del rumore"),
    ("integrative", "Misure preventive e protettive integrative rispetto al PSC"),
    ("procedure", "Procedure complementari e di dettaglio richieste dal PSC"),
    ("telefoni", "Telefoni e indirizzi utili locali (ospedale, Vigili del fuoco, Carabinieri, ASL...)"),
    ("emergenze", "Gestione delle emergenze: primo soccorso, antincendio ed evacuazione"),
]


def chiave_testo(impresa_id: int, sezione: str) -> str:
    return f"pos:{impresa_id}:{sezione}"


def elenco_dpi(testi: list[str]) -> list[str]:
    """Unisce gli elenchi di DPI delle schede di rischio eliminando i doppioni."""
    visti: dict[str, str] = {}
    for testo in testi:
        for voce in re.split(r",|;|\n", testo):
            voce = voce.strip().rstrip(".")
            if voce:
                visti.setdefault(voce.lower(), voce[0].upper() + voce[1:])
    return sorted(visti.values(), key=str.lower)


def _periodo(impresa: dict, periodo_fasi: tuple | None) -> str:
    """Date dei lavori indicate per l'impresa o, in mancanza, ricavate dalle sue lavorazioni."""
    inizio = analisi.data(impresa["data_inizio"]) or (periodo_fasi[0] if periodo_fasi else None)
    fine = analisi.data(impresa["data_fine"]) or (periodo_fasi[1] if periodo_fasi else None)
    if not (inizio and fine):
        return ""
    return f"dal {data_it(inizio)} al {data_it(fine)} ({(fine - inizio).days + 1} giorni)"


# Numeri di emergenza nazionali (il numero unico 112 è attivo in tutta Italia).
NUMERI_EMERGENZA = [
    ("Numero unico di emergenza", "112"),
    ("Carabinieri", "112"),
    ("Polizia di Stato", "113"),
    ("Vigili del fuoco", "115"),
    ("Emergenza sanitaria", "118"),
]

# Documenti da tenere in cantiere (testo redatto per SicurCantiere dagli obblighi del D.Lgs. 81/2008
# e delle norme collegate; le voci «ove previsto» dipendono dal cantiere).
DOCUMENTI_IN_CANTIERE = [
    "copia della notifica preliminare, esposta in modo visibile (art. 99)",
    "piano di sicurezza e di coordinamento e fascicolo dell'opera, se previsti (artt. 100 e 91)",
    "piano operativo di sicurezza di ciascuna impresa esecutrice e relativi aggiornamenti (art. 96)",
    "titolo abilitativo dei lavori",
    "iscrizione alla Camera di commercio, DURC e documentazione dell'idoneità tecnico-professionale "
    "delle imprese e dei lavoratori autonomi (art. 90 e Allegato XVII)",
    "libro unico del lavoro o documentazione equivalente e tessere di riconoscimento dei lavoratori (art. 26 c. 8)",
    "attestati di formazione, addestramento e abilitazione all'uso delle attrezzature (artt. 37 e 73)",
    "giudizi di idoneità alla mansione rilasciati dal medico competente (art. 41)",
    "verbali di consegna dei DPI (art. 77)",
    "schede di dati di sicurezza delle sostanze e miscele pericolose usate (Reg. CE 1907/2006)",
    "libretti d'uso e manutenzione e dichiarazioni di conformità CE di macchine e attrezzature (art. 70)",
    "verbali delle verifiche periodiche delle attrezzature di sollevamento e registro delle verifiche "
    "trimestrali di funi e catene (art. 71 e Allegato VII), ove previsto",
    "PiMUS, autorizzazione ministeriale e, se richiesto, progetto del ponteggio (artt. 133 e 136), ove previsto",
    "dichiarazione di conformità dell'impianto elettrico e di terra del cantiere (D.M. 37/2008) e relativa "
    "denuncia (D.P.R. 462/2001), ove prevista",
    "piano di lavoro per la rimozione dell'amianto (art. 256), ove previsto",
    "verbali dei sopralluoghi e delle riunioni di coordinamento con il coordinatore per l'esecuzione",
]


def _foto_cantiere(repo: Repository) -> list[str]:
    return [r for r in repo.testo("foto_cantiere").splitlines() if r.strip()]


def _tabella_rischi(doc, rischi: list[dict]) -> None:
    righe = []
    for r in rischi:
        livello = analisi.livello_rischio(r["probabilita"], r["danno"])
        descrizione = f"{r['descrizione']} ({r['fonte']})" if r["fonte"] else r["descrizione"]
        righe.append([descrizione, f"{livello} - {analisi.classe_rischio(livello)}", r["misure"]])
    if righe:
        t = stile.tabella_colorata(doc, ["Rischio", "Livello", "Misure preventive e protettive"], righe, 8.5,
                                   larghezze=[5.2, 2.4, 9.4])
        for riga, r in zip(t.rows[1:], rischi):
            livello = analisi.livello_rischio(r["probabilita"], r["danno"])
            colore = "F4B6B6" if livello >= 9 else "F9D3A6" if livello >= 6 else "FFF2B3" if livello >= 3 else "D5EFD8"
            stile.ombreggia_cella(riga.cells[1], colore)


def _esito_breve(repo: Repository, valutazione: dict) -> list[str]:
    """Indice e classe della valutazione, o «da eseguire» se i parametri mancano."""
    if da_eseguire(repo, valutazione["id"]):
        return ["-", "Valutazione da eseguire"]
    e = METODI[valutazione["metodo"]].esegui(repo.dati_valutazione(valutazione["id"]))
    return [e.indice_testo, e.classe]


def _icone_dpi(doc, testi: list[str], etichetta: str) -> None:
    icone = segnaletica.icone_dpi(testi)
    if not icone:
        return
    p = doc.add_paragraph()
    p.add_run(f"{etichetta}  ").bold = True
    stile.icone_in_riga(p, [segnaletica.SEGNALI[i].file for i in icone], 0.85)


def genera_pos(repo: Repository, impresa_id: int, destinazione: str | Path) -> Path:
    impresa = repo.leggi("imprese", impresa_id)
    if impresa is None:
        raise ValueError("Impresa inesistente")
    fasi = repo.elenco("fasi", impresa_id=impresa_id)
    lavoratori = repo.elenco("lavoratori", impresa_id=impresa_id)
    rischi = [r for f in fasi for r in repo.rischi(fase_id=f["id"])]
    attrezzature = [a for f in fasi for a in repo.elenco("attrezzature", fase_id=f["id"])]
    c = repo.cantiere()
    soggetti = anagrafiche.soggetti(repo)
    committente = next((s for s in soggetti if s["ruolo"] == anagrafiche.COMMITTENTE), None)
    periodo = analisi.periodo_lavori(fasi)
    sezioni = {chiave: repo.testo(chiave_testo(impresa_id, chiave)) for chiave, _ in SEZIONI_POS}
    figure = [{"ruolo": anagrafiche.DATORE, "nome": impresa["datore_lavoro"], "qualifica": "Datore di lavoro",
               "telefono": impresa["telefono"], "email": impresa["email"], "codice_fiscale": "",
               "formazione": ""}] + anagrafiche.figure(repo, impresa)

    doc = Document()
    stile.imposta_stili(doc)
    comune = f"{c['comune']} ({c['provincia']})" if c["provincia"] else c["comune"]
    stile.imposta_pagina(doc, " - ".join(v for v in (c["nome"], comune) if v))

    # --- copertina ---
    rspp = next((f for f in figure if f["ruolo"] == anagrafiche.RSPP and f["nome"]), None)
    firme = [("L'impresa - il datore di lavoro", f"{impresa['ragione_sociale']} - {impresa['datore_lavoro']}".strip(" -"))]
    if rspp and rspp["nome"] != impresa["datore_lavoro"]:
        firme.append(("Il responsabile del servizio di prevenzione e protezione", rspp["nome"]))
    stile.copertina(
        doc,
        comune=f"Comune di {c['comune']}" if c["comune"] else "",
        titolo="PIANO OPERATIVO DI SICUREZZA",
        riferimento="art. 89 c. 1 lett. h), art. 96 e Allegato XV punto 3.2 del D.Lgs. 9 aprile 2008, n. 81",
        voci=[("Oggetto", c["nome"] or c["descrizione"]),
              ("Committente", committente["nome"] if committente else ""),
              ("Cantiere", indirizzo_cantiere(c))],
        data_luogo=f"{c['comune'] or ''}, {data_it(date.today())}".strip(", "),
        firme=firme,
        impresa=[impresa["ragione_sociale"], anagrafiche.sede(impresa),
                 " - ".join(v for v in (f"Tel. {impresa['telefono']}" if impresa["telefono"] else "",
                                        impresa["pec"] or impresa["email"]) if v)],
        logo=impresa["logo"],
    )

    # --- indice e revisioni ---
    stile.indice(doc)
    revisioni = repo.elenco("revisioni", documento=f"POS:{impresa_id}")
    if revisioni:
        stile.banda(doc, "Revisioni del documento")
        stile.tabella_colorata(doc, ["Rev.", "Data", "Descrizione"], [
            [str(r["numero"]), data_it(r["data"]), r["descrizione"]] for r in revisioni])

    # --- a) dati generali ---
    stile.capitolo(doc, "Lavoro", "punto 3.2.1, lettera a), punto 1, Allegato XV del D.Lgs. 81/2008")
    stile.banda(doc, "Caratteristiche generali dell'opera")
    stile.scheda(doc, [
        ("Natura dell'opera", c["natura_opera"]),
        ("Oggetto", c["nome"]),
        ("Titolo abilitativo", c["titolo_abilitativo"]),
        ("Importo presunto dei lavori", euro(c["importo_lavori"]) if c["importo_lavori"] else ""),
        ("Data di inizio dei lavori", data_it(periodo[0]) if periodo else ""),
        ("Data di fine dei lavori (presunta)", data_it(periodo[1]) if periodo else ""),
    ])
    if c["descrizione"].strip():
        stile.banda(doc, "Descrizione sintetica dell'opera")
        testo_o_segnaposto(doc, c["descrizione"])
    stile.banda(doc, "Dati del cantiere")
    stile.scheda(doc, [("Indirizzo", c["indirizzo"]), ("CAP", c["cap"]), ("Città", comune)])
    for n, foto in enumerate(_foto_cantiere(repo), 1):
        stile.immagine(doc, foto, 11, f"Foto {n} - area del cantiere")

    stile.capitolo(doc, "Committente e soggetti con compiti di sicurezza",
                   "punto 3.2.1, lettera a), Allegato XV del D.Lgs. 81/2008")
    for s in soggetti:
        stile.banda(doc, s["ruolo"])
        if s.get("nota"):
            stile.nota(doc, s["nota"])
        stile.scheda(doc, [
            ("Nominativo / ragione sociale", s["nome"]), ("Nella persona di", s["rappresentante"]),
            ("Qualifica", s["qualifica"]), ("Indirizzo", s["indirizzo"]), ("Telefono", s["telefono"]),
            ("Email", s["email"]), ("PEC", s["pec"]), ("Codice fiscale", s["codice_fiscale"]),
            ("Partita IVA", s["partita_iva"]),
        ])
    if not soggetti:
        testo_o_segnaposto(doc, "")

    stile.capitolo(doc, "Dati identificativi dell'impresa esecutrice",
                   "punto 3.2.1, lettera a), punto 1, Allegato XV del D.Lgs. 81/2008")
    affidataria = repo.leggi("imprese", impresa["affidataria_id"]) if impresa["affidataria_id"] else None
    stile.banda(doc, "Dati dell'impresa")
    stile.scheda(doc, [
        ("Tipologia nel cantiere", impresa["tipo"]),
        ("Ragione sociale", impresa["ragione_sociale"]),
        ("Forma giuridica", impresa["forma_giuridica"]),
        ("Datore di lavoro", impresa["datore_lavoro"]),
        ("Sede legale", anagrafiche.sede(impresa)),
        ("Telefono / fax", " - ".join(v for v in (impresa["telefono"], impresa["fax"]) if v)),
        ("Email / PEC", " - ".join(v for v in (impresa["email"], impresa["pec"]) if v)),
        ("Ufficio o recapito in cantiere", impresa["recapito_cantiere"]),
        ("Codice fiscale", impresa["codice_fiscale"]),
        ("Partita IVA", impresa["partita_iva"]),
        ("Registro imprese / CCIAA", impresa["cciaa"]),
        ("Codice ATECO", impresa["ateco"]),
        ("Posizione INAIL", impresa["inail"]),
        ("Matricola INPS", impresa["inps"]),
        ("Cassa edile", impresa["cassa_edile"]),
        ("Contratto collettivo applicato", impresa["ccnl"]),
        ("Organico medio annuo", impresa["organico"]),
        ("DURC valido fino al", data_it(impresa["durc_scadenza"]) if impresa["durc_scadenza"] else ""),
        ("Documento di valutazione dei rischi del", data_it(impresa["dvr_data"]) if impresa["dvr_data"] else ""),
        ("Attestazione SOA", impresa["soa"]),
        ("Certificazioni", impresa["certificazioni"]),
    ])
    stile.banda(doc, "Dati dei lavori")
    stile.scheda(doc, [
        ("Impresa affidataria", affidataria["ragione_sociale"] if affidataria else ""),
        ("Autorizzazione al subappalto", impresa["aut_subappalto"]),
        ("Tipologia dei lavori affidati", impresa["tipologia_lavori"]),
        ("Importo dei lavori da eseguire", euro(impresa["importo_lavori"]) if impresa["importo_lavori"] else ""),
        ("Inizio e fine dei lavori dell'impresa", _periodo(impresa, periodo)),
        ("Orario di lavoro", impresa["orario_lavoro"]),
        ("Soggetto incaricato (art. 97)", " - ".join(v for v in (impresa["incaricato_97"],
                                                                 impresa["incaricato_97_mansione"]) if v)),
    ])
    subaffidatari = repo.elenco("imprese", affidataria_id=impresa_id)
    if subaffidatari:
        stile.banda(doc, "Imprese subappaltatrici e lavoratori autonomi")
        stile.tabella_colorata(doc, ["Ragione sociale", "Tipologia", "P. IVA", "Lavori affidati"], [
            [i["ragione_sociale"], i["tipo"], i["partita_iva"], i["tipologia_lavori"]] for i in subaffidatari])

    stile.capitolo(doc, "Responsabili e figure della sicurezza",
                   "punto 3.2.1, lettera a), punti 3, 4, 5 e 6, Allegato XV del D.Lgs. 81/2008")
    for f in figure:
        stile.banda(doc, f["ruolo"])
        if f.get("nota"):
            stile.nota(doc, f["nota"])
        stile.scheda(doc, [("Nome e cognome", f["nome"]), ("Qualifica", f["qualifica"]),
                           ("Telefono", f["telefono"]), ("Email", f.get("email", "")),
                           ("Codice fiscale", f.get("codice_fiscale", "")), ("Formazione", f["formazione"])])

    stile.capitolo(doc, "Numero e relative qualifiche dei lavoratori dipendenti e autonomi",
                   "punto 3.2.1, lettera a), punto 7, Allegato XV del D.Lgs. 81/2008")
    dipendenti = [l for l in lavoratori if not l["autonomo"]]
    autonomi = [l for l in lavoratori if l["autonomo"]]
    if dipendenti:
        stile.banda(doc, f"Lavoratori dipendenti: {len(dipendenti)}")
        stile.tabella_colorata(doc, ["N.", "Qualifica", "Cognome", "Nome", "Matricola", "Idoneità fino al"], [
            [str(n), l["mansione"], l["cognome"], l["nome"], l["matricola"],
             data_it(l["idoneita_scadenza"]) if l["idoneita_scadenza"] else ""]
            for n, l in enumerate(dipendenti, 1)])
    if autonomi:
        stile.banda(doc, f"Lavoratori autonomi: {len(autonomi)}")
        stile.tabella_colorata(doc, ["N.", "Qualifica", "Nominativo", "C.F. / P. IVA", "Periodo"], [
            [str(n), l["mansione"], anagrafiche.nominativo(l), l["partita_iva"] or l["codice_fiscale"],
             f"{data_it(l['data_inizio'])} - {data_it(l['data_fine'])}"] for n, l in enumerate(autonomi, 1)])
    if not lavoratori:
        testo_o_segnaposto(doc, "")

    stile.capitolo(doc, "Specifiche mansioni inerenti la sicurezza",
                   "punto 3.2.1, lettera b), Allegato XV del D.Lgs. 81/2008")
    for ruolo in dict.fromkeys(f["ruolo"] for f in figure):
        nomi = ", ".join(f["nome"] for f in figure if f["ruolo"] == ruolo and f["nome"])
        stile.banda(doc, f"{ruolo}{f' - {nomi}' if nomi else ''}")
        p = doc.add_paragraph()
        r = p.add_run("Specifiche mansioni inerenti la sicurezza: ")
        r.font.color.rgb = stile.BLU
        p.add_run(anagrafiche.compiti(ruolo) or "compiti secondo l'incarico ricevuto.")
        nota_ruolo = next((f.get("nota") for f in figure if f["ruolo"] == ruolo and f.get("nota")), None)
        if nota_ruolo:
            p.add_run(" " + nota_ruolo).italic = True

    stile.capitolo(doc, "Specifiche attività e singole lavorazioni svolte in cantiere",
                   "punto 3.2.1, lettera a), punto 2, Allegato XV del D.Lgs. 81/2008")
    lavorazione_corrente = None
    for f in fasi:
        if f["lavorazione"] and f["lavorazione"] != lavorazione_corrente:
            lavorazione_corrente = f["lavorazione"]
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.add_run(f"- {lavorazione_corrente}").bold = True
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8 if f["lavorazione"] else 0)
        p.paragraph_format.space_after = Pt(0)
        p.add_run(f"- {nome_fase(f)}").bold = not f["lavorazione"]
    if not fasi:
        testo_o_segnaposto(doc, "")
    stile.nota(doc, f"Se non diversamente specificato, le attività e le lavorazioni indicate sono svolte "
                    f"dall'impresa esecutrice «{impresa['ragione_sociale']}».")

    # --- documentazione ---
    stile.capitolo(doc, "Documentazione")
    stile.banda(doc, "Telefoni e indirizzi utili")
    stile.scheda(doc, [(nome, f"tel. {numero}") for nome, numero in NUMERI_EMERGENZA])
    if sezioni["telefoni"].strip():
        testo_o_segnaposto(doc, sezioni["telefoni"])
    stile.banda(doc, "Documentazione da custodire in cantiere")
    doc.add_paragraph("Presso il cantiere o gli uffici di cantiere l'impresa tiene a disposizione degli "
                      "organi di vigilanza e del coordinatore la seguente documentazione:")
    for voce in DOCUMENTI_IN_CANTIERE:
        doc.add_paragraph(voce, style="List Bullet")

    # --- d) opere provvisionali, macchine, impianti ---
    stile.capitolo(doc, "Elenco delle opere provvisionali, macchine e impianti utilizzati in cantiere",
                   "punto 3.2.1, lettera d), Allegato XV del D.Lgs. 81/2008")
    per_tipo: dict[str, list[str]] = {}
    for a in attrezzature:
        per_tipo.setdefault(a["tipo"], [])
        if a["nome"] not in per_tipo[a["tipo"]]:
            per_tipo[a["tipo"]].append(a["nome"])
    if per_tipo:
        doc.add_paragraph().add_run("AD USO ESCLUSIVO DELL'IMPRESA").bold = True
        for tipo in ("Macchina", "Attrezzo", "Opera provvisionale"):
            if tipo in per_tipo:
                p = doc.add_paragraph()
                p.add_run({"Macchina": "Macchine:", "Attrezzo": "Attrezzi:",
                           "Opera provvisionale": "Opere provvisionali:"}[tipo]).bold = True
                for n, nome in enumerate(sorted(per_tipo[tipo]), 1):
                    q = doc.add_paragraph(f"{n})  {nome};")
                    q.paragraph_format.left_indent = Cm(0.6)
                    q.paragraph_format.space_after = Pt(0)
    else:
        testo_o_segnaposto(doc, "")

    # --- e) sostanze ---
    stile.capitolo(doc, "Elenco delle sostanze e miscele pericolose",
                   "punto 3.2.1, lettera e), Allegato XV del D.Lgs. 81/2008")
    doc.add_heading("Sostanze e miscele pericolose e relative schede di sicurezza", 2)
    aggiungi_sostanze(doc, repo, sostanze_impresa(repo, impresa_id), livello=3)
    if sezioni["sostanze"].strip():
        testo_o_segnaposto(doc, sezioni["sostanze"])

    # --- f) rumore e vibrazioni ---
    valutazioni = valutazioni_impresa(repo, impresa_id)
    for metodo, titolo in (("rumore", dict(SEZIONI_POS)["rumore"]),
                           ("vibrazioni", "Esito della valutazione delle vibrazioni")):
        elenco = [v for v in valutazioni if v["metodo"] == metodo]
        if metodo == "vibrazioni" and not elenco:
            continue
        stile.capitolo(doc, titolo, "punto 3.2.1, lettera f), Allegato XV del D.Lgs. 81/2008" if metodo == "rumore"
                       else "Titolo VIII, Capo III del D.Lgs. 81/2008")
        if elenco:
            stile.tabella_colorata(doc, ["Mansione / gruppo omogeneo", "Esposizione", "Esito"], [
                [v["mansione"] or v["titolo"], *_esito_breve(repo, v)] for v in elenco])
            stile.nota(doc, "Il dettaglio dei calcoli e delle misure è riportato nel capitolo delle "
                            "valutazioni dei rischi specifici.")
        if metodo == "rumore" and (sezioni["rumore"].strip() or not elenco):
            testo_o_segnaposto(doc, sezioni["rumore"])

    # --- segnaletica ---
    stile.capitolo(doc, "Segnaletica prevista nel cantiere", "Titolo V e Allegati XXIV-XXXII del D.Lgs. 81/2008")
    stile.cartello_testo(doc, "Vietato l'accesso ai non addetti ai lavori")
    doc.add_paragraph()
    cartelli = segnaletica.segnali_cantiere([r["descrizione"] for r in rischi],
                                            [r["dpi"] for r in rischi] + [a["dpi"] for a in attrezzature])
    stile.griglia_immagini(doc, [(s.file, s.testo) for s in cartelli], colonne=4, lato_cm=2.8)
    stile.nota(doc, "Rappresentazione indicativa: i cartelli posati in cantiere sono conformi alla UNI EN ISO 7010 "
                    "e collocati in prossimità dei pericoli e degli accessi.")

    # --- c), g), h) attività, misure, procedure per lavorazione e fase ---
    stile.capitolo(doc, "Descrizione delle attività di cantiere, delle modalità organizzative e dei turni di lavoro",
                   "punto 3.2.1, lettere c), g) e h), Allegato XV del D.Lgs. 81/2008")
    testo_o_segnaposto(doc, sezioni["attivita"])
    lavorazione_corrente = None
    for f in fasi:
        if f["lavorazione"] != lavorazione_corrente:
            lavorazione_corrente = f["lavorazione"]
            if lavorazione_corrente:
                stile.riquadro(doc, lavorazione_corrente)
                gruppo = [g for g in fasi if g["lavorazione"] == lavorazione_corrente]
                p = doc.add_paragraph()
                p.add_run("La lavorazione è suddivisa nelle seguenti fasi: ").bold = True
                p.add_run("; ".join(nome_fase(g) for g in gruppo) + ".")
        doc.add_heading(nome_fase(f) + (" (fase)" if f["lavorazione"] else ""), 3)
        if f["dettaglio"]:
            doc.add_paragraph(f["dettaglio"])
        struttura = struttura_fase(repo, f["id"])
        impegnati = repo.lavoratori_fase(f["id"])
        stile.scheda(doc, [
            ("Macchine utilizzate", ", ".join(a["nome"] for a in struttura["macchine"])),
            ("Attrezzi e opere provvisionali", ", ".join(a["nome"] for a in struttura["attrezzi"])),
            ("Addetto", f["addetto"]),
            ("Lavoratori impegnati", ", ".join(f"{anagrafiche.nominativo(l)} ({l['mansione']})" if l["mansione"]
                                               else anagrafiche.nominativo(l) for l in impegnati)),
            ("Modalità organizzative", f["modalita"]),
            ("Turni di lavoro", f["turni"]),
            ("Procedure complementari richieste dal PSC", f["procedure_psc"]),
        ], 5.2)
        rischi_fase = [r for r in rischi if r["fase_id"] == f["id"]]
        _icone_dpi(doc, [r["dpi"] for r in struttura["rischi_lavorazione"]], "DPI dell'addetto:")
        esclusi = repo.rischi_esclusi(fase_id=f["id"])
        if rischi_fase:
            _tabella_rischi(doc, rischi_fase)
            collegate = [v for v in valutazioni if f["id"] in repo.fasi_valutazione(v["id"])]
            if collegate:
                stile.nota(doc, "Rischi specifici valutati con metodo dedicato (capitolo «Valutazioni dei rischi "
                                "specifici»): " + "; ".join(
                                    f"{v['titolo'] or METODI[v['metodo']].nome} - {_esito_breve(repo, v)[1]}"
                                    for v in collegate if v["metodo"] in METODI) + ".")
        else:
            testo_o_segnaposto(doc, "")
        if esclusi:
            stile.banda(doc, "Rischi valutati come non presenti")
            stile.tabella_colorata(doc, ["Rischio", "Motivazione"], [
                [f"{r['descrizione']} ({r['fonte']})" if r["fonte"] else r["descrizione"], r["motivazione"] or "-"]
                for r in esclusi], 8.5, colore_intestazione="6B7B8C", larghezze=[6, 11])

    for chiave, lettera in (("integrative", "g"), ("procedure", "h")):
        stile.capitolo(doc, dict(SEZIONI_POS)[chiave], f"punto 3.2.1, lettera {lettera}), Allegato XV del D.Lgs. 81/2008")
        testo_o_segnaposto(doc, sezioni[chiave])

    # --- valutazioni specifiche ---
    if valutazioni:
        stile.capitolo(doc, "Valutazioni dei rischi specifici", "Titoli VI, VIII e IX del D.Lgs. 81/2008")
        stile.tabella_colorata(doc, ["Valutazione", "Lavorazioni", "Metodo", "Indice", "Classe"], [
            [v["titolo"] or METODI[v["metodo"]].nome, lavorazioni_valutazione(repo, v) or "-",
             METODI[v["metodo"]].norma, *_esito_breve(repo, v)]
            for v in valutazioni if v["metodo"] in METODI], 8.5, larghezze=[3.6, 4.6, 3.4, 2.2, 3.2])
        for v in valutazioni:
            aggiungi_valutazione(doc, repo, v, livello=2)

    # --- schede di attrezzature e macchine ---
    for titolo, macchine in (("Attrezzature utilizzate nelle lavorazioni", False),
                             ("Macchine utilizzate nelle lavorazioni", True)):
        elenco = {a["nome"]: a for a in attrezzature if (a["tipo"] == MACCHINA) == macchine}
        if not elenco:
            continue
        stile.capitolo(doc, titolo, "artt. 70-73 e Allegati V-VII del D.Lgs. 81/2008")
        p = doc.add_paragraph()
        p.add_run("Elenco: ").bold = True
        p.add_run("; ".join(sorted(elenco)) + ".")
        for nome, a in sorted(elenco.items()):
            stile.riquadro(doc, nome, a["tipo"].lower(), "7F6000", "FFF2CC")
            if a["descrizione"]:
                doc.add_paragraph(a["descrizione"])
            propri = list(dict.fromkeys(r["descrizione"] for r in rischi if r["fonte"] == nome))
            if propri:
                doc.add_paragraph().add_run(f"Rischi generati dall'uso {'della macchina' if macchine else 'dell attrezzo'}:"
                                            .replace("dell attrezzo", "dell'attrezzo")).bold = True
                for n, voce in enumerate(propri, 1):
                    q = doc.add_paragraph(f"{n})  {voce};")
                    q.paragraph_format.left_indent = Cm(0.6)
                    q.paragraph_format.space_after = Pt(0)
            if a["misure"]:
                doc.add_paragraph().add_run("Misure preventive e protettive:").bold = True
                for riga in a["misure"].split("\n"):
                    doc.add_paragraph(riga)
            if a["dpi"]:
                p = doc.add_paragraph()
                chi = "dell'operatore" if macchine else "dell'utilizzatore"
                p.add_run(f"DPI {chi}: ").bold = True
                p.add_run(a["dpi"] + ".")
                stile.icone_in_riga(doc.add_paragraph(),
                                    [segnaletica.SEGNALI[i].file for i in segnaletica.icone_dpi(a["dpi"])], 0.8)

    # --- i) DPI ---
    stile.capitolo(doc, "Elenco dei dispositivi di protezione individuale forniti ai lavoratori",
                   "punto 3.2.1, lettera i), Allegato XV del D.Lgs. 81/2008")
    con_dpi = [l for l in lavoratori if anagrafiche.dpi_lavoratore(l)]
    if con_dpi:
        stile.tabella_colorata(doc, ["Qualifica", "Lavoratore", "DPI forniti"], [
            [l["mansione"], anagrafiche.nominativo(l), "; ".join(anagrafiche.dpi_lavoratore(l))] for l in con_dpi])
    dpi = elenco_dpi([r["dpi"] for r in rischi])
    if dpi:
        stile.banda(doc, "DPI previsti dalle schede di rischio delle lavorazioni")
        stile.griglia_immagini(doc, [
            (segnaletica.SEGNALI[i].file, segnaletica.SEGNALI[i].testo.replace(" obbligatori", "")
             .replace(" obbligatoria", "")) for i in segnaletica.icone_dpi(dpi)], colonne=5, lato_cm=2.0, dimensione_testo=6.5)
        for voce in dpi:
            doc.add_paragraph(voce, style="List Bullet")
    elif not con_dpi:
        testo_o_segnaposto(doc, "")

    # --- l) formazione ---
    stile.capitolo(doc, "Formazione e informazione dei lavoratori",
                   "punto 3.2.1, lettera l), Allegato XV del D.Lgs. 81/2008")
    doc.add_paragraph(
        "I lavoratori ricevono informazione, formazione e addestramento ai sensi degli artt. 36, 37 e 77 "
        "del D.Lgs. 81/2008 e dell'Accordo Stato-Regioni vigente; la documentazione è conservata in cantiere.")
    if lavoratori:
        stile.tabella_colorata(doc, ["N.", "Lavoratore", "Qualifica", "Informazione, formazione e addestramento"], [
            [str(n), anagrafiche.nominativo(l), l["mansione"], "; ".join(anagrafiche.formazione_lavoratore(l)) or "-"]
            for n, l in enumerate(lavoratori, 1)])

    # --- emergenze ---
    stile.capitolo(doc, "Organizzazione del servizio di primo soccorso, antincendio ed evacuazione dei lavoratori",
                   "artt. 18, 43-46 e punto 3.2.1, lettera a), punto 3, Allegato XV del D.Lgs. 81/2008")
    addetti = [(f["ruolo"], f["nome"]) for f in figure if f["ruolo"].startswith("Addetto") and f["nome"]]
    if addetti:
        stile.scheda(doc, addetti)
    testo_o_segnaposto(doc, sezioni["emergenze"])
    stile.griglia_immagini(doc, [(segnaletica.SEGNALI[k].file, segnaletica.SEGNALI[k].testo)
                                 for k in ("soccorso", "estintore")], colonne=2, lato_cm=2.4)

    # --- conclusioni ---
    stile.capitolo(doc, "Conclusioni generali")
    doc.add_paragraph("Al presente piano operativo di sicurezza sono allegati, e ne fanno parte integrante:")
    allegati = ["schede di dati di sicurezza delle sostanze e miscele pericolose;",
                "documentazione sulla formazione e informazione dei lavoratori;",
                "verbali di consegna dei DPI;"]
    if valutazioni:
        allegati.insert(0, "relazioni delle valutazioni dei rischi specifici;")
    stile.elenco_lettere(doc, allegati)
    doc.add_paragraph()
    doc.add_paragraph(f"{c['comune'] or ''}, {data_it(date.today())}".strip(", "))
    firma(doc, "Il Datore di lavoro")
    doc.add_paragraph("Per presa visione, il Rappresentante dei lavoratori per la sicurezza")
    doc.add_paragraph("_______________________________")

    destinazione = Path(destinazione)
    doc.save(destinazione)
    return destinazione
