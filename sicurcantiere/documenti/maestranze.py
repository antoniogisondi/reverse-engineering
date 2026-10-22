"""Tesserini di riconoscimento e schede dei lavoratori con mansioni, rischi, DPI e formazione."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt

from sicurcantiere import analisi, anagrafiche
from sicurcantiere.documenti.comune import (
    data_it,
    indirizzo_cantiere,
    nome_fase,
    nuovo_documento,
    ombreggia,
    salva,
    tabella,
    tabella_voci,
)
from sicurcantiere.documenti.pos import elenco_dpi
from sicurcantiere.repository import Repository
from sicurcantiere.valutazioni import METODI


def lavoratori_impresa(repo: Repository, impresa_id: int) -> list[dict[str, Any]]:
    """Dipendenti prima, poi lavoratori autonomi."""
    return sorted(repo.elenco("lavoratori", impresa_id=impresa_id), key=lambda l: (l["autonomo"], l["id"]))


def _immagine(run, percorso: str, larghezza: Cm) -> bool:
    if not percorso or not Path(percorso).is_file():
        return False
    try:
        run.add_picture(percorso, width=larghezza)
        return True
    except Exception:  # noqa: BLE001 - immagine non leggibile: il tesserino resta senza foto
        return False


def genera_tesserini(repo: Repository, impresa_id: int, destinazione: str | Path) -> Path:
    """Tessere di riconoscimento (art. 18, comma 1, lettera u), art. 21, comma 1, lettera c), e art. 26, comma 8,
    del D.Lgs. 81/2008; art. 5 della L. 136/2010)."""
    impresa = repo.leggi("imprese", impresa_id)
    if impresa is None:
        raise ValueError("Impresa inesistente")
    lavoratori = lavoratori_impresa(repo, impresa_id) or [{"nome": "", "cognome": "", "autonomo": 0}]
    doc = nuovo_documento()
    for sezione in doc.sections:
        sezione.top_margin = sezione.bottom_margin = Cm(1.5)
    doc.add_paragraph().add_run(f"Tessere di riconoscimento - {impresa['ragione_sociale']}").bold = True
    doc.add_paragraph("Ritagliare lungo i bordi e inserire nell'apposita custodia. Il lavoratore espone la tessera in "
                      "modo visibile per tutta la permanenza in cantiere.")

    griglia = doc.add_table(rows=0, cols=2)
    griglia.style = "Table Grid"
    griglia.alignment = WD_TABLE_ALIGNMENT.CENTER
    for n, l in enumerate(lavoratori):
        if n % 2 == 0:
            celle = griglia.add_row().cells
        cella = celle[n % 2]
        cella.width = Cm(8.6)
        interna = cella.add_table(rows=1, cols=2)
        foto, dati = interna.rows[0].cells
        foto.width, dati.width = Cm(2.6), Cm(5.6)
        if not _immagine(foto.paragraphs[0].add_run(), l.get("foto", ""), Cm(2.3)):
            foto.paragraphs[0].add_run("\nFOTO\n\n").font.size = Pt(7)
        cella.paragraphs[0].add_run(impresa["ragione_sociale"] or "Impresa").bold = True
        ombreggia(cella, "F2F2F2")
        righe = [("", anagrafiche.nominativo(l) or "________________")]
        if l.get("data_nascita"):
            righe.append(("Nato/a il", f"{data_it(l['data_nascita'])} {l.get('luogo_nascita', '')}".strip()))
        if l.get("autonomo"):
            righe.append(("", "Lavoratore autonomo"))
            if l.get("partita_iva") or l.get("codice_fiscale"):
                righe.append(("P. IVA / C.F.", l.get("partita_iva") or l.get("codice_fiscale")))
        else:
            righe.append(("Qualifica", l.get("mansione", "")))
            if l.get("matricola"):
                righe.append(("Matricola", l["matricola"]))
            righe.append(("Assunto il", data_it(l.get("data_assunzione")) if l.get("data_assunzione") else "______"))
            righe.append(("Datore di lavoro", impresa["datore_lavoro"]))
        if impresa["aut_subappalto"]:
            righe.append(("Aut. subappalto", impresa["aut_subappalto"]))
        dati.paragraphs[0].text = ""
        for i, (etichetta, valore) in enumerate(righe):
            p = dati.paragraphs[0] if i == 0 else dati.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            if etichetta:
                p.add_run(f"{etichetta}: ").font.size = Pt(7)
            r = p.add_run(valore or "-")
            r.font.size = Pt(9 if i == 0 else 7.5)
            r.bold = i == 0
        nota = cella.add_paragraph("Tessera di riconoscimento ai sensi degli artt. 18, 21 e 26 del D.Lgs. 81/2008")
        nota.runs[0].font.size = Pt(6)
    return salva(doc, destinazione)


def mansioni_lavoratore(repo: Repository, lavoratore: dict[str, Any]) -> list[dict[str, Any]]:
    """Per ogni lavorazione assegnata: fase, attrezzature usate e rischi a cui il lavoratore è esposto.

    Se per una lavorazione sono indicate le attrezzature usate, i rischi delle altre attrezzature
    della fase non si attribuiscono al lavoratore.
    """
    coppie = repo.mansioni(lavoratore["id"])
    risultato = []
    for fase_id in dict.fromkeys(f for f, _ in sorted(coppie, key=lambda x: (x[0], x[1] or 0))):
        fase = repo.leggi("fasi", fase_id)
        if fase is None:
            continue
        attrezzature = {a["id"]: a for a in repo.elenco("attrezzature", fase_id=fase_id)}
        usate = [attrezzature[a] for f, a in coppie if f == fase_id and a in attrezzature]
        nomi_fase = {a["nome"] for a in attrezzature.values()}
        nomi_usate = {a["nome"] for a in usate}
        rischi = [r for r in repo.rischi(fase_id=fase_id)
                  if not usate or r["fonte"] not in nomi_fase or r["fonte"] in nomi_usate]
        risultato.append({"fase": fase, "attrezzature": usate or list(attrezzature.values()), "rischi": rischi})
    return risultato


def genera_schede_lavoratori(repo: Repository, impresa_id: int, destinazione: str | Path) -> Path:
    """Una scheda per lavoratore: dati, mansioni, macchine e attrezzature, rischi, DPI, formazione."""
    impresa = repo.leggi("imprese", impresa_id)
    if impresa is None:
        raise ValueError("Impresa inesistente")
    c = repo.cantiere()
    valutazioni = [v for v in repo.elenco("valutazioni") if v["fase_id"] is not None and v["metodo"] in METODI]
    doc = nuovo_documento()
    lavoratori = lavoratori_impresa(repo, impresa_id)
    if not lavoratori:
        doc.add_paragraph("Nessun lavoratore indicato per l'impresa.")
    for n, l in enumerate(lavoratori):
        if n:
            doc.add_page_break()
        doc.add_heading(f"SCHEDA DEL LAVORATORE - {anagrafiche.nominativo(l) or 'nominativo da indicare'}", 1)
        doc.add_paragraph(f"Cantiere: {c['nome']} - {indirizzo_cantiere(c)}".strip(" -"))
        voci = [
            ("Impresa", impresa["ragione_sociale"]),
            ("Posizione", "Lavoratore autonomo" if l["autonomo"] else "Lavoratore dipendente"),
            ("Qualifica / mansione", l["mansione"]),
            ("Data e luogo di nascita",
             f"{data_it(l['data_nascita']) if l['data_nascita'] else ''} {l['luogo_nascita']}".strip()),
        ]
        if l["autonomo"]:
            voci += [("Codice fiscale / P. IVA", " - ".join(v for v in (l["codice_fiscale"], l["partita_iva"]) if v)),
                     ("Periodo in cantiere", f"{data_it(l['data_inizio'])} - {data_it(l['data_fine'])}")]
        else:
            voci += [("Matricola", l["matricola"]),
                     ("Data di assunzione", data_it(l["data_assunzione"]) if l["data_assunzione"] else "")]
        voci.append(("Idoneità sanitaria fino al", data_it(l["idoneita_scadenza"]) if l["idoneita_scadenza"] else ""))
        tabella_voci(doc, voci)

        mansioni = mansioni_lavoratore(repo, l)
        doc.add_heading("Lavorazioni svolte e macchine o attrezzature utilizzate", 2)
        if mansioni:
            tabella(doc, ["Lavorazione", "Periodo", "Macchine e attrezzature"], [
                [nome_fase(m["fase"]), f"{data_it(m['fase']['data_inizio'])} - {data_it(m['fase']['data_fine'])}",
                 ", ".join(a["nome"] for a in m["attrezzature"]) or "-"] for m in mansioni
            ])
        else:
            doc.add_paragraph().add_run("Mansioni non ancora assegnate.").italic = True

        rischi = [r for m in mansioni for r in m["rischi"]]
        doc.add_heading("Rischi cui è esposto il lavoratore e misure di prevenzione", 2)
        if rischi:
            visti: dict[str, dict] = {}
            for r in rischi:
                chiave = r["descrizione"].lower()
                livello = analisi.livello_rischio(r["probabilita"], r["danno"])
                if chiave not in visti or livello > visti[chiave]["livello"]:
                    visti[chiave] = {**r, "livello": livello}
            tabella(doc, ["Rischio", "Livello", "Misure"], [
                [r["descrizione"], f"{r['livello']} {analisi.classe_rischio(r['livello'])}", r["misure"]]
                for r in sorted(visti.values(), key=lambda x: -x["livello"])
            ], dimensione=8.5)
        else:
            doc.add_paragraph("-")

        fasi_lavoratore = {m["fase"]["id"] for m in mansioni}
        sue = [v for v in valutazioni if v["fase_id"] in fasi_lavoratore]
        if sue:
            doc.add_heading("Valutazioni dei rischi specifici", 2)
            tabella(doc, ["Valutazione", "Metodo", "Esito"], [
                [v["titolo"] or METODI[v["metodo"]].nome, METODI[v["metodo"]].norma,
                 f"{(e := METODI[v['metodo']].esegui(repo.dati_valutazione(v['id']))).indice_testo} - {e.classe}"]
                for v in sue
            ])

        doc.add_heading("Dispositivi di protezione individuale", 2)
        dpi = anagrafiche.dpi_lavoratore(l) or elenco_dpi([r["dpi"] for r in rischi])
        doc.add_paragraph("; ".join(dpi) if dpi else "-")
        doc.add_heading("Informazione, formazione e addestramento", 2)
        formazione = anagrafiche.formazione_lavoratore(l)
        for voce in formazione or ["-"]:
            doc.add_paragraph(voce, style="List Bullet")

        doc.add_paragraph()
        doc.add_paragraph("Il lavoratore dichiara di aver ricevuto copia della presente scheda e le informazioni "
                          "sui rischi e sulle misure relative alle proprie mansioni.")
        doc.add_paragraph("Il datore di lavoro                                         Il lavoratore")
        doc.add_paragraph("_____________________________            _____________________________")
    return salva(doc, destinazione)
