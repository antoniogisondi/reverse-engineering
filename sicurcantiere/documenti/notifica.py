"""Notifica preliminare (art. 99 e Allegato XII del D.Lgs. 81/2008).

Tutti i campi sono ricavati dai dati del progetto; quelli mancanti restano evidenziati nella bozza.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document

from sicurcantiere import anagrafiche, analisi
from sicurcantiere.documenti import stile
from sicurcantiere.documenti.comune import data_it, euro, indirizzo_cantiere, salva
from sicurcantiere.documenti.composizione import CONSIGLIATO, OBBLIGATORIO, Manca
from sicurcantiere.repository import Repository

CANTIERE, SOGGETTI, IMPRESE, FASI = "Dati generali", "Soggetti", "Imprese", "Fasi"


def _soggetto(soggetti: list[dict], ruolo: str) -> str:
    trovati = [s for s in soggetti if s["ruolo"] == ruolo and s["nome"].strip()]
    return "; ".join(
        ", ".join(v for v in (s["nome"], s["codice_fiscale"], s["indirizzo"]) if v) for s in trovati
    )


def _o(valore: str, dato: str, sezione: str, livello: str = OBBLIGATORIO):
    return valore if str(valore or "").strip() else Manca(dato, sezione, livello)


def genera_notifica(repo: Repository, destinazione: str | Path) -> Path:
    with stile.colori_documento(*stile.COLORI_PSC):
        return _genera(repo, destinazione)


def _genera(repo: Repository, destinazione: str | Path) -> Path:
    c = repo.cantiere()
    soggetti = anagrafiche.soggetti(repo)
    imprese = repo.elenco("imprese")
    fasi = repo.elenco("fasi")
    periodo = analisi.periodo_lavori(fasi)
    durata = analisi.giorni_lavorativi(*periodo) if periodo else 0
    comune = f"{c['comune']} ({c['provincia']})" if c["provincia"] else c["comune"]

    doc = Document()
    stile.imposta_stili(doc)
    stile.imposta_pagina(doc, " · ".join(v for v in (c["nome"], comune) if v),
                         (f"Notifica preliminare · {c['nome']}".strip(" ·"), data_it(date.today())))
    for sezione in doc.sections:
        sezione.different_first_page_header_footer = False  # modulo senza copertina

    p = doc.add_paragraph()
    stile._spaziatura(p, 0, 2)
    stile._run(p, "D.LGS. 9 APRILE 2008, N. 81 · ART. 99 E ALLEGATO XII", 7.5, stile.ACCENTO, grassetto=True)
    p = doc.add_paragraph()
    stile._spaziatura(p, 0, 4)
    stile._bordo_paragrafo(p, ("bottom",), stile.ACCENTO, 18, 6)
    stile._run(p, "Notifica preliminare", 24, stile.INCHIOSTRO, carattere=stile.CARATTERE_TITOLI)
    p = doc.add_paragraph()
    stile._spaziatura(p, 4, 12)
    stile._run(p, "All'Azienda sanitaria locale e all'Ispettorato territoriale del lavoro competenti per territorio "
                  "e, per le opere pubbliche, al Prefetto.", 9, stile.TESTO2)

    stile.scheda(doc, [
        ("1. Data della comunicazione", data_it(date.today())),
        ("2. Indirizzo del cantiere", _o(indirizzo_cantiere(c), "indirizzo del cantiere", CANTIERE)),
        ("3. Committente", _o(_soggetto(soggetti, anagrafiche.COMMITTENTE), "committente (nome, codice fiscale, "
                                                                          "indirizzo)", SOGGETTI)),
        ("4. Natura dell'opera", _o(c["natura_opera"] or c["nome"], "natura dell'opera", CANTIERE)),
        ("5. Responsabile dei lavori", _o(_soggetto(soggetti, anagrafiche.RESPONSABILE_LAVORI),
                                          "responsabile dei lavori, se nominato", SOGGETTI, CONSIGLIATO)),
        ("6. Coordinatore per la progettazione", _o(_soggetto(soggetti, anagrafiche.CSP),
                                                    "coordinatore per la progettazione", SOGGETTI)),
        ("7. Coordinatore per l'esecuzione", _o(_soggetto(soggetti, anagrafiche.CSE),
                                                "coordinatore per l'esecuzione", SOGGETTI)),
        ("8. Data presunta d'inizio dei lavori", _o(data_it(periodo[0]) if periodo else "",
                                                    "data di inizio dei lavori (date delle fasi)", FASI)),
        ("9. Durata presunta dei lavori", _o(f"{durata} giorni lavorativi" if durata else "",
                                             "durata dei lavori (date delle fasi)", FASI)),
        ("10. Numero massimo presunto dei lavoratori in cantiere",
         _o(str(analisi.picco_lavoratori(fasi)) if fasi else "", "numero massimo di lavoratori (fasi)", FASI)),
        ("11. Numero previsto di imprese e di lavoratori autonomi",
         _o(str(len(imprese)) if imprese else "", "numero di imprese e lavoratori autonomi", IMPRESE)),
        ("13. Ammontare complessivo presunto dei lavori",
         _o(euro(c["importo_lavori"]) if c["importo_lavori"] else "", "importo dei lavori", CANTIERE)),
    ], 6.2)

    stile.banda(doc, "12. Imprese e lavoratori autonomi già selezionati")
    if imprese:
        stile.tabella_colorata(doc, ["Ragione sociale", "Tipo", "P. IVA / C.F.", "Sede"], [
            [i["ragione_sociale"], i["tipo"], i["partita_iva"] or i["codice_fiscale"], anagrafiche.sede(i)]
            for i in imprese], larghezze=[5.2, 3.0, 3.4, 5.8])
    else:
        stile.testo(doc, "", "imprese già selezionate, se presenti", IMPRESE, CONSIGLIATO)

    p = doc.add_paragraph(f"{c['comune'] or ''}, {data_it(date.today())}".strip(", "))
    stile._spaziatura(p, 14, 8)
    responsabile = next((s for s in soggetti if s["ruolo"] in (anagrafiche.RESPONSABILE_LAVORI,
                                                              anagrafiche.COMMITTENTE) and s["nome"].strip()), None)
    stile.firme(doc, [("Il committente o il responsabile dei lavori", responsabile["nome"] if responsabile else "")])
    stile.nota(doc, "La notifica è trasmessa prima dell'inizio dei lavori e ne va esposta una copia in cantiere, in "
                    "maniera visibile, custodita a disposizione dell'organo di vigilanza (art. 99). Va aggiornata in "
                    "caso di variazioni.")
    return salva(doc, destinazione)
